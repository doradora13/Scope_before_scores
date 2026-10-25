# -*- coding: utf-8 -*-
"""
23_mycovenant7_extract.py - analysis-ready extract of the MyCovenant 7th release (January 2026) and the data checks of
Supplementary Section S17 and Table S24.

Inputs (set in lib.py; a missing file is downloaded on first use):
  MYCOV7_DF1, MYCOV7_DF2, MYCOV7_DF3  MyCovenant 7th release (JRC Data Catalogue, doi:10.2905/JRC.MDWNXYR; CC BY 4.0)
  DATA_PATH, JRC_XLSX                 2021 extract; JRC source tables (Table 2: direct and indirect emissions by sector;
                                      Table 3: NUTS3 codes and ancillary covariates)
  LAU_GEOJSON, NUTS16_GEOJSON         Eurostat GISCO LAU 2018 and NUTS 2016 level-3 boundaries
                                      ((c) EuroGeographics for the administrative boundaries)
  NUTS16_CODES, NUTS21_CODES          NUTS 2016 and 2021 code lists
  HDD_CSV, GDP_CSV                    Eurostat nrg_chddr2_a (heating degree days) and nama_10r_3gdp (GDP; used only for a check)
Steps
  1. Release sheets to CSV (cached in OUT_DIR/work/mycov7_csv); inventories in the schema of the 2021 extract: reporting
     sectors mapped to the six 2021 sectors (Table S24, Panel A), zero-emission sectors treated as unreported, one inventory
     per signatory, type and year (latest entry; the earliest entry for a sensitivity check), match with the 2021 extract.
  2. Covariates: 2021 ancillary values for 2021 signatories. For new signatories: population and area of the matched LAU 2018
     unit (name match within country, else point in polygon), its NUTS 2016 region, NUTS3 GDP per capita from the 2021 table,
     heating degree days from the 2021 table for the same region and reference year, else Eurostat nrg_chddr2_a.
  3. Checks reported in Section S17 and Table S24: sector mapping on inventories with the same id in both releases, NUTS codes,
     external covariates validated on 2021 signatories, the same inventories in both analysis samples and the per-capita screen.
Outputs: work/MyCovenant7_SFA_ready.csv and work/MyCovenant7_SFA_ready_earliest.csv (compared with the copies at
  MYCOV7_DATA_PATH and MYCOV7_EARLIEST_PATH when present), tables/MyCovenant7_Processing.xlsx,
  tables/mycovenant7_processing_numbers.csv, tables/check_mycovenant7_processing.csv (comparison with the numbers printed in
  the manuscript), tables/Table_S24.tex (LaTeX for Table S24).
Requires shapely 2.0 or later (pip install shapely) for the point-in-polygon matching. About 5 minutes (first run: also the
downloads and the conversion of the release workbooks).
"""
import csv
import re
import json
import hashlib
import itertools
import unicodedata
import warnings
from lib import *
warnings.filterwarnings("ignore")
try:
    from shapely.geometry import shape, Point
    from shapely import STRtree
except ImportError:
    raise SystemExit("Stage 23 needs shapely 2.0 or later: pip install shapely")
import openpyxl

SEC6 = SECTOR_NAMES
SC = SECTOR_COLS
# MyCovenant aggregated_sector -> sector of the 2021 dataset (None: not used, as in the 2021 definition)
MAP_MAIN = {"Municipal buildings, equipment/facilities": "municipal",
            "Tertiary (non municipal) buildings, equipment/facilities": "tertiary",
            "Residential buildings": "residential",
            "Agriculture, Forestry, Fisheries (excluding AFOLU)": "residential",
            "Buildings, equipment/facilities non allocated": "residential",
            "Non-ETS": "manufacturing", "Industry Non-ETS": "manufacturing",
            "Transportation": "transport", "Waste/wastewater": "waste",
            "ETS (not recommended)": None, "Industry-ETS": None, "Other": None}
EXTRA = {"Agriculture, Forestry, Fisheries (excluding AFOLU)": "agri", "Buildings, equipment/facilities non allocated": "nonalloc",
         "ETS (not recommended)": "ets", "Industry-ETS": "ets", "Other": "other", "Residential buildings": "resid_only"}
CC = {"gr": "el", "gb": "uk"}           # country codes of the 2021 extract
TOL = 1e-4                              # 'identical' values: relative difference of at most 0.01%
# SHA-256 of the source files used for the manuscript; a different hash means a different vintage of that file
SHA256 = {"df1_Signatories_7th Release - January 2026.xlsx": "85713538d96cd3a35ee52ab3403b52dadd2e06cace1ea8ddb0dbed2b7f4f980d",
          "df2_Action_Plans_7th Release - January 2026.xlsx": "993081d842c5f36d296ee1f972f88f963d86a7cb5e0329960d5e0661009a256a",
          "df3_Monitoring_Reports_7th Release - January 2026.xlsx": "c0a35e9ba5b0fb8c8e845f087beef8b331dab659e568cd513c397ef12520f4d5",
          "LAU_RG_01M_2018_4326.geojson": "f0ae30114987bbc81937170ef3475434d0b344019987ea912717409194902cf3",
          "NUTS_RG_01M_2016_4326_LEVL_3.geojson": "8a1cc246592a20656abb403fbfc9451d5abc607fb3b583a65501bf45128725f7",
          "NUTS_AT_2016.csv": "eb2ebbe550d08702d623d2454bfaa6e4e9f976f6bd1ad89a4e6d7ccbfa436cd2",
          "NUTS_AT_2021.csv": "a44b9d59972afa9e827371623b8375a9cb3eef07606e86ffe98d8941491ad22c",
          "nrg_chddr2_a.csv": "c117af1e61cc15fc5c752085d71e5a24db8c84dd6a6b6573d3de73f9eecbc241",
          "nama_10r_3gdp.csv": "f6aeeeb73e85a6d1d2dfb712fd977bfd5821ea52c66b9fd5289c74da9e9f92be",
          "GlobalCoM.xlsx": "ddf055dbdb066f613ef40c4ece92c43bf22c9b0ff9cd71f749daac64856c1d89",
          "GlobalCoM_SFA_ready.csv": "308d0963de0a73ff92d95ade01eccc95b0a2d719ec17af8bd3cac3e94b5bb7aa"}

banner("23 MyCovenant 7th release: analysis-ready extract and data checks (Supplementary Section S17, Table S24)")
nums = []


def note(name, value, txt=""):
    nums.append(dict(name=name, value=value, note=txt))


# ------------------------------------------------------------------ source files
src = [(MYCOV7_DF1, MYCOV7_DF1_URL), (MYCOV7_DF2, MYCOV7_DF2_URL), (MYCOV7_DF3, MYCOV7_DF3_URL), (LAU_GEOJSON, LAU_GEOJSON_URL),
       (NUTS16_GEOJSON, NUTS16_GEOJSON_URL), (NUTS16_CODES, NUTS16_CODES_URL), (NUTS21_CODES, NUTS21_CODES_URL),
       (HDD_CSV, HDD_CSV_URL), (GDP_CSV, GDP_CSV_URL), (JRC_XLSX, JRC_XLSX_URL)]
for p, u in src:
    ensure_file(p, u)
hashes = []
for p in [s[0] for s in src] + [DATA_PATH]:
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    exp = SHA256.get(os.path.basename(p))
    hashes.append(dict(file=os.path.basename(p), sha256=h.hexdigest(), same_as_manuscript=(exp == h.hexdigest()) if exp else None))
hashes = pd.DataFrame(hashes)
print(hashes[["file", "same_as_manuscript"]].to_string(index=False))
if not hashes.same_as_manuscript.fillna(True).all():
    print("NOTE: a source file differs from the vintage used for the manuscript; results can differ slightly.")


# ------------------------------------------------------------------ 1. release sheets -> CSV (read once)
CSV_DIR = wpath("mycov7_csv")
os.makedirs(CSV_DIR, exist_ok=True)


def sheet_csv(xlsx, sheet, name):
    """Write one worksheet to OUT_DIR/work/mycov7_csv/<name>.csv (all cells as written by the csv module) and return its path."""
    out = os.path.join(CSV_DIR, name + ".csv")
    if os.path.exists(out):
        return out
    print(f"reading sheet {sheet} of {os.path.basename(xlsx)} ...", flush=True)
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    with open(out + ".part", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        for row in wb[sheet].iter_rows(values_only=True):
            w.writerow(["" if v is None else v for v in row])
    wb.close()
    os.replace(out + ".part", out)
    return out


sig = pd.read_csv(sheet_csv(MYCOV7_DF1, "1", "df1_1"), low_memory=False, dtype={"organisation_id": str, "gcom_id": str, "nuts3_code": str})
ap = pd.read_csv(sheet_csv(MYCOV7_DF2, "2", "df2_2"), low_memory=False, dtype={"organisation_id": str, "action_plan_id": str})
e2 = pd.read_csv(sheet_csv(MYCOV7_DF2, "2D_1", "df2_2D_1"), low_memory=False, dtype={"action_plan_id": str, "emission_inventory_id": str})
mr = pd.read_csv(sheet_csv(MYCOV7_DF3, "3", "df3_3"), low_memory=False, dtype={"organisation_id": str, "action_plan_id": str, "monitoring_report_id": str})
e3 = pd.read_csv(sheet_csv(MYCOV7_DF3, "3D_1", "df3_3D_1"), low_memory=False, dtype={"monitoring_report_id": str, "emission_inventory_id": str})

# ------------------------------------------------------------------ inventories
e2["src"], e3["src"] = "AP", "MR"
e2 = e2.merge(ap[["action_plan_id", "organisation_id", "approval_date", "action_plan_pillar_valid_mitigation", "number_of_members"]],
              on="action_plan_id", how="left")
e3 = e3.merge(mr[["monitoring_report_id", "organisation_id", "action_plan_id", "monitoring_report_type", "monitoring_report_submission_date",
                  "monitoring_report_resubmission_date", "monitoring_report_pillar_valid_mitigation"]], on="monitoring_report_id", how="left")
e = pd.concat([e2, e3], ignore_index=True)
e["sec"] = e.aggregated_sector.map(MAP_MAIN)
unknown = set(e.aggregated_sector) - set(MAP_MAIN)
assert not unknown, f"unknown reporting sectors: {unknown}"
e["ind"] = e.aggregated_carrier.isin(["Electricity", "District heating and cooling"])
n_ap, n_mr = e2.emission_inventory_id.nunique(), e3.emission_inventory_id.nunique()
print(f"rows {len(e)}; inventories {e.emission_inventory_id.nunique()} (action plans {n_ap}, monitoring reports {n_mr})")
note("inventories_action_plans", n_ap, "S17: inventories in action plans")
note("inventories_monitoring_reports", n_mr, "S17: inventories in monitoring reports")

key = "emission_inventory_id"
meta = e.groupby(key).agg(src=("src", "first"), organisation_id=("organisation_id", "first"), action_plan_id=("action_plan_id", "first"),
                          monitoring_report_id=("monitoring_report_id", "first"), flag=("baseline_year_flag", "first"),
                          inventory_year=("inventory_year", "first"), population_inventory_year=("population_in_the_inventory_year", "first"),
                          emission_factor_type=("emission_factor_type", "first"), emission_unit=("emission_reporting_unit", "first"),
                          ap_valid=("action_plan_pillar_valid_mitigation", "first"), mr_valid=("monitoring_report_pillar_valid_mitigation", "first"),
                          mr_type=("monitoring_report_type", "first"), n_members=("number_of_members", "first"))
em = e[e.sec.notna()].pivot_table(index=key, columns="sec", values="emission_measure", aggfunc="sum")
ac = e[e.sec.notna()].pivot_table(index=key, columns="sec", values="energy_measure", aggfunc="sum")
ex = e[e.aggregated_sector.isin(EXTRA)].assign(x=lambda d: d.aggregated_sector.map(EXTRA)).pivot_table(index=key, columns="x", values="emission_measure", aggfunc="sum")
en = e[e.sec.notna() & (e.sec != "waste")]
di = en.pivot_table(index=key, columns="ind", values=["emission_measure", "energy_measure"], aggfunc="sum")
di.columns = [f"{'emissions' if a == 'emission_measure' else 'activity_data'}_{'ind' if b else 'dir'}" for a, b in di.columns]
inv = meta.join(em.reindex(columns=SEC6).add_suffix("_emissions_tonnes")).join(ac.reindex(columns=SEC6[:5]).add_suffix("_activity_MWh"))
inv = inv.join(ex.add_prefix("x_")).join(di).reset_index()
# a sector is reported when its emissions are positive (zero sums are template placeholders)
nz = {c: int((inv[c] == 0).sum()) for c in SC}
for c in SC:
    inv.loc[inv[c] <= 0, c] = np.nan
print("zero-emission sectors treated as not reported:", nz)
note("zero_emission_sector_entries", sum(nz.values()), "S17 and Table S24: sector entries with zero emissions treated as unreported")
for s in SEC6[:5]:
    inv.loc[inv[f"{s}_emissions_tonnes"].isna(), f"{s}_activity_MWh"] = np.nan
inv["inventory_type"] = np.where(inv.flag.astype(str) == "True", "baseline_emission_inventory", "monitoring_emission_inventory")
inv["total_emissions_reported_tonnes"] = inv[SC].sum(axis=1, min_count=1)
inv["total_activity_MWh"] = inv[[f"{s}_activity_MWh" for s in SEC6[:5]]].sum(axis=1, min_count=1)
inv["reported_sector_count"] = inv[SC].notna().sum(axis=1)
inv["core5_complete_flag"] = (inv[SC[:5]].notna().sum(axis=1) == 5).astype(int)
inv["energy_emissions_tonnes"] = inv[SC[:5]].sum(axis=1, min_count=1)
# alternative mapping (Table S25): residential = 'Residential buildings' only
inv["residential_raw_emissions_tonnes"] = inv["x_resid_only"].where(inv["x_resid_only"] > 0)
sg = sig.drop_duplicates("organisation_id").set_index("organisation_id")
inv["GCoM_ID"] = inv.organisation_id.map(sg.gcom_id)
inv["signatory_name"] = inv.organisation_id.map(sg.organisation_name)
inv["country_code"] = inv.organisation_id.map(sg.country_code).replace(CC)
inv["release_lat"] = inv.organisation_id.map(sg.organisation_latitude)
inv["release_lon"] = inv.organisation_id.map(sg.organisation_longitude)
inv["release_nuts3_2021"] = inv.organisation_id.map(sg.nuts3_code)
inv["adhesion_type"] = inv.organisation_id.map(sg.adhesion_type)
inv["eid_num"] = pd.to_numeric(inv.emission_inventory_id.str.replace("_med", "", regex=False), errors="coerce")
# joint plans of several municipalities (number_of_members > 0) and the baseline year declared in each action plan
joint_orgs = set(ap.loc[ap.number_of_members > 0, "organisation_id"])
joint_ids = set(sig.loc[sig.organisation_id.isin(joint_orgs), "gcom_id"].dropna())
inv["joint_plan"] = inv.GCoM_ID.isin(joint_ids).astype(int)
plan_bei = e2[e2.baseline_year_flag.astype(str) == "True"].groupby("action_plan_id").inventory_year.max()
inv["plan_bei_year"] = inv.action_plan_id.astype(str).map(plan_bei)
n0 = len(inv)
inv = inv[inv.GCoM_ID.notna() & (inv.total_emissions_reported_tonnes > 0)].copy()
print(f"inventories with a signatory id and a positive six-sector total: {len(inv)} of {n0}")
note("inventories_with_id_and_positive_total", len(inv))

# ------------------------------------------------------------------ one inventory per (signatory, type, year)
inv = inv.sort_values(["GCoM_ID", "inventory_type", "inventory_year", "eid_num"])
grp = inv.groupby(["GCoM_ID", "inventory_type", "inventory_year"])
inv["n_versions"] = grp.emission_inventory_id.transform("size")
multi = grp.size() > 1
tot_diff = grp.total_emissions_reported_tonnes.agg(lambda s: s.max() / s.min() - 1 > 1e-6)
print(f"keys with more than one inventory: {int(multi.sum())} (inventories {int((inv.n_versions > 1).sum())}); "
      f"of these keys, {int((tot_diff & multi).sum())} have different totals")
note("keys_several_entries", int(multi.sum()), "S17 and Table S24")
note("keys_several_entries_different_totals", int((tot_diff & multi).sum()), "S17 and Table S24")
versions = {"latest": inv.drop_duplicates(["GCoM_ID", "inventory_type", "inventory_year"], keep="last").copy(),     # highest inventory id
            "earliest": inv.drop_duplicates(["GCoM_ID", "inventory_type", "inventory_year"], keep="first").copy()}
inv = versions["latest"]
print(f"after de-duplication: {len(inv)} inventories; signatories {inv.GCoM_ID.nunique()}; BEI {int((inv.inventory_type == 'baseline_emission_inventory').sum())}, "
      f"MEI {int((inv.inventory_type == 'monitoring_emission_inventory').sum())}")
note("distinct_keys", len(inv), "S17 and Table S24: one inventory per signatory, type and year")

# ------------------------------------------------------------------ match with the 2021 extract
old = pd.read_csv(DATA_PATH)
k21 = set(zip(old.GCoM_ID, old.inventory_type, old.inventory_year))
id21 = set(old.emission_inventory_id.astype(str))
for v in versions.values():
    v["in2021_key"] = [k in k21 for k in zip(v.GCoM_ID, v.inventory_type, v.inventory_year)]
    v["in2021_id"] = v.emission_inventory_id.isin(id21)
    v["sig_in2021"] = v.GCoM_ID.isin(set(old.GCoM_ID))
    v["new_inventory"] = ~v.in2021_key
found = len(k21 & set(zip(inv.GCoM_ID, inv.inventory_type, inv.inventory_year)))
print(f"2021 inventories found by signatory, type and year: {found} of {len(k21)}")
note("inventories_2021_source", len(k21)); note("inventories_2021_matched_by_key", found)

# ------------------------------------------------------------------ 2. geography and covariates
_, t3 = load_jrc()
t3 = t3.drop_duplicates("GCoM_ID").set_index("GCoM_ID")
t3 = t3.rename(columns={"GDP per capita at NUTS3 \n[Euro per inhabitant]": "gdp", "Heating Degree-Days (HDD)": "hdd", "reference year HDD": "hddy",
                        "population in 2018": "pop18", "CITY or Greater City": "cgc", "degree of urbanisation": "degurba",
                        "GHG emissions per capita in GCoM sectors_EDGAR\n[tCO2-eq/year]": "edgar_pc"})
gdp_by_region = t3.dropna(subset=["NUTS3", "gdp"]).groupby("NUTS3").gdp.first()
hdd_by_region_year = t3.dropna(subset=["NUTS3", "hdd", "hddy"]).groupby(["NUTS3", "hddy"]).hdd.first()
hdd_all = pd.read_csv(HDD_CSV, dtype={"geo": str})
hes = hdd_all[(hdd_all.indic_nrg == "HDD") & (hdd_all.geo.str.len() == 5)].set_index(["geo", "TIME_PERIOD"]).OBS_VALUE

sigs = inv.groupby("GCoM_ID").agg(name=("signatory_name", "first"), cty=("country_code", "first"), lat=("release_lat", "first"),
                                   lon=("release_lon", "first"), nuts21=("release_nuts3_2021", "first"))
first_bei = inv[inv.inventory_type == "baseline_emission_inventory"].groupby("GCoM_ID").inventory_year.min()
sigs["ref_year"] = first_bei.reindex(sigs.index).fillna(inv.groupby("GCoM_ID").inventory_year.min())
sigs["in_t3"] = sigs.index.isin(t3.index)


def norm(s):
    """Signatory or LAU name for matching: ASCII, lower case, without common prefixes such as 'comune di'."""
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower()
    s = re.sub(r"[^a-z0-9/ ]+", " ", s)
    s = re.sub(r"\b(comune di|municipality of|city of|ayuntamiento de|ajuntament de|concello de|ville de|commune de|gemeente|stadt|gemeinde|"
               r"municipio de|camara municipal de|obcina|opcina|grad|miasto|gmina)\b", " ", s)
    return re.sub(r"\s+", " ", s).strip()


print("reading the GISCO boundaries ...", flush=True)
with open(LAU_GEOJSON, encoding="utf-8") as f:
    lau = json.load(f)["features"]
lau_geom = [shape(f["geometry"]) for f in lau]
lau_p = pd.DataFrame([f["properties"] for f in lau])
lau_p["cty"] = lau_p.CNTR_CODE.str.lower()
lau_p["rep"] = [g.representative_point() for g in lau_geom]
lau_p["rx"] = [p.x for p in lau_p.rep]; lau_p["ry"] = [p.y for p in lau_p.rep]
name_idx = {}
for i, (c, nm) in enumerate(zip(lau_p.cty, lau_p.LAU_NAME)):
    for part in str(nm).split("/"):
        name_idx.setdefault((c, norm(part)), []).append(i)
tree = STRtree(lau_geom)
with open(NUTS16_GEOJSON, encoding="utf-8") as f:
    n16 = json.load(f)["features"]
n16_geom = [shape(f["geometry"]) for f in n16]; n16_id = [f["properties"]["NUTS_ID"] for f in n16]
t16 = STRtree(n16_geom)


def pip_one(tr, geoms, p, maxd=0.05):
    """Index of the polygon containing p; otherwise the nearest polygon within maxd degrees (about 5 km), else None."""
    hit = tr.query(p, predicate="within")
    if len(hit):
        return int(hit[0])
    j, d = tr.query_nearest(p, return_distance=True)
    return int(j[0]) if len(j) and d[0] < maxd else None


def lau_match(name, cty, lat, lon):
    """Name match within country (nearest to the release coordinates if several); otherwise point in polygon."""
    p = Point(lon, lat) if np.isfinite(lat) and np.isfinite(lon) else None
    cand = name_idx.get((cty, norm(name)), [])
    if cand:
        if len(cand) == 1 or p is None:
            return cand[0], "name"
        d = [(lau_p.rx[i] - lon) ** 2 + (lau_p.ry[i] - lat) ** 2 for i in cand]
        return cand[int(np.argmin(d))], "name+nearest"
    if p is not None:
        j = pip_one(tree, lau_geom, p)
        if j is not None and lau_p.cty[j] == cty:
            return j, "pip"
    return None, "none"


rows = []
for gid, r in sigs.iterrows():
    j, how = lau_match(r["name"], r.cty, r.lat, r.lon)
    out = dict(GCoM_ID=gid, lau_how=how)
    if j is not None:
        out.update(lau_id=lau_p.GISCO_ID[j], lau_name=lau_p.LAU_NAME[j], pop18_lau=lau_p.POP_2018[j], area_lau=lau_p.AREA_KM2[j])
        k = pip_one(t16, n16_geom, lau_p.rep[j])
        out["nuts3_lau16"] = n16_id[k] if k is not None else None
    rows.append(out)
geo = pd.DataFrame(rows).set_index("GCoM_ID")
# LAU 2018 units below the municipal level: Portugal (freguesias) and Greece (local communities) are aggregated to the
# municipality by the first four code digits; Ireland (electoral divisions) and Cyprus are left missing
lau_p["muni_key"] = np.where(lau_p.cty.isin(["pt", "el"]), lau_p.cty + "_" + lau_p.LAU_ID.astype(str).str[:4], lau_p.GISCO_ID)
agg = lau_p.groupby("muni_key").agg(pop_m=("POP_2018", "sum"), area_m=("AREA_KM2", "sum"))
key_of = dict(zip(lau_p.GISCO_ID, lau_p.muni_key))
geo["muni_key"] = geo.lau_id.map(key_of)
sub = geo.lau_id.notna() & geo.lau_id.str[:2].isin(["PT", "EL"])
geo.loc[sub, "pop18_lau"] = geo.loc[sub, "muni_key"].map(agg.pop_m)
geo.loc[sub, "area_lau"] = geo.loc[sub, "muni_key"].map(agg.area_m)
sub_na = geo.lau_id.notna() & geo.lau_id.str[:2].isin(["IE", "CY"])
geo.loc[sub_na, ["pop18_lau", "area_lau"]] = np.nan
sigs = sigs.join(geo)
L16 = set(pd.read_csv(NUTS16_CODES, dtype=str).NUTS_ID)
L21 = set(pd.read_csv(NUTS21_CODES, dtype=str).NUTS_ID)
pt16 = []
for lat, lon in zip(sigs.lat, sigs.lon):
    k = pip_one(t16, n16_geom, Point(lon, lat)) if np.isfinite(lat) and np.isfinite(lon) else None
    pt16.append(n16_id[k] if k is not None else None)
sigs["nuts3_pip16"] = pt16
# NUTS3 (2016): region of the matched LAU unit, else the release code when valid in NUTS 2016, else point in polygon
sigs["nuts3_ext"] = sigs.nuts3_lau16.where(sigs.nuts3_lau16.notna(), np.where(sigs.nuts21.isin(L16), sigs.nuts21, sigs.nuts3_pip16))

# validation of the external geography on signatories of the 2021 table
v = sigs[sigs.in_t3].join(t3[["NUTS3", "LAU2", "pop18", "area"]])
okn = v.NUTS3.notna()
lr_dens = np.log((v.pop18_lau / v.area_lau) / (v.pop18 / v.area))
agree = float(np.mean(v.nuts3_ext[okn] == v.NUTS3[okn]))
print(f"validation on {len(v)} signatories of the 2021 table: LAU match {v.lau_how.value_counts().to_dict()}; NUTS3 (2016) agreement {agree:.4f}; "
      f"density median |log ratio| {np.nanmedian(np.abs(lr_dens)):.4f}")
note("validation_nuts3_agreement", agree, "S17 and Table S24"); note("validation_density_median_abs_log_ratio", float(np.nanmedian(np.abs(lr_dens))))

# covariates: 2021 table for its signatories; otherwise external
sigs["nuts3"] = np.where(sigs.in_t3, t3.NUTS3.reindex(sigs.index), sigs.nuts3_ext)
sigs["gdp_per_capita_nuts3_eur"] = np.where(sigs.in_t3, t3.gdp.reindex(sigs.index), sigs.nuts3_ext.map(gdp_by_region))
hdd_ext, hdd_src = [], []
for gid, r in sigs.iterrows():
    if r.in_t3:
        hdd_ext.append(t3.hdd.get(gid, np.nan)); hdd_src.append("jrc2021"); continue
    y = int(r.ref_year) if np.isfinite(r.ref_year) else None
    val, srcname = np.nan, "none"
    if y is not None and r.nuts3_ext is not None:
        val = hdd_by_region_year.get((r.nuts3_ext, y), np.nan); srcname = "jrc2021_region_year"
        if not np.isfinite(val):
            for code in [r.nuts3_ext, r.nuts21]:
                val = hes.get((code, max(min(y, 2025), 1980)), np.nan)
                if np.isfinite(val):
                    srcname = "eurostat"; break
    hdd_ext.append(val); hdd_src.append(srcname if np.isfinite(val) else "none")
sigs["hdd"] = hdd_ext; sigs["hdd_source"] = hdd_src
sigs["population_2018"] = np.where(sigs.in_t3, t3.pop18.reindex(sigs.index), sigs.pop18_lau)
sigs["area"] = np.where(sigs.in_t3, t3.area.reindex(sigs.index), sigs.area_lau)
sigs["latitude"] = np.where(sigs.in_t3, t3.latitude.reindex(sigs.index), sigs.lat)
sigs["longitude"] = np.where(sigs.in_t3, t3.longitude.reindex(sigs.index), sigs.lon)
sigs["city_or_greater_city"] = np.where(sigs.in_t3, t3.cgc.reindex(sigs.index), np.nan)
sigs["covariate_source"] = np.where(sigs.in_t3, "jrc2021", "external")
cov_cols = ["nuts3", "gdp_per_capita_nuts3_eur", "hdd", "hdd_source", "population_2018", "area", "latitude", "longitude",
            "city_or_greater_city", "covariate_source", "lau_how", "ref_year"]
need = ["population_inventory_year", "gdp_per_capita_nuts3_eur", "population_density_2018_per_km2", "hdd"]
for k_, v_ in versions.items():
    v_ = v_.join(sigs[cov_cols], on="GCoM_ID")
    v_["NUTS3"] = v_.nuts3
    v_["population_density_2018_per_km2"] = v_.population_2018 / v_.area
    v_["reported_emissions_per_capita"] = v_.total_emissions_reported_tonnes / v_.population_inventory_year
    v_.loc[~(v_.population_inventory_year > 0), "population_inventory_year"] = np.nan
    v_["emission_unit"] = v_.emission_unit.replace({"tonnes of CO2/year": "tonnes CO2", "tonnes of CO2-eq/year": "tonnes CO2 equivalent"})
    v_["covars_ok7"] = v_[need].notna().all(axis=1) & (v_.population_density_2018_per_km2 > 0)
    versions[k_] = v_
inv = versions["latest"]
print(f"covariates complete: {int(inv.covars_ok7.sum())} of {len(inv)}; heating degree days by source: {inv.hdd_source.value_counts().to_dict()}")

# ------------------------------------------------------------------ write the extracts and compare with the copies set in lib.py
for k_, path in [("latest", MYCOV7_DATA_PATH), ("earliest", MYCOV7_EARLIEST_PATH)]:
    out = wpath("MyCovenant7_SFA_ready" + ("" if k_ == "latest" else "_earliest") + ".csv")
    versions[k_].to_csv(out, index=False)
    if os.path.exists(path) and os.path.abspath(path) != os.path.abspath(out):
        a = pd.read_csv(out, **RELEASE_DTYPES); b = pd.read_csv(path, **RELEASE_DTYPES)
        same = a.shape == b.shape and list(a.columns) == list(b.columns) and a.equals(b)
        if not same and a.shape == b.shape and list(a.columns) == list(b.columns):
            num = a.select_dtypes("number").columns
            same = bool(np.allclose(a[num].values, b[num].values, rtol=1e-10, atol=0, equal_nan=True)) and a.drop(columns=num).equals(b.drop(columns=num))
        print(f"{os.path.basename(out)}: {'identical to' if same else 'DIFFERENT from'} {path}")
        note(f"extract_{k_}_same_as_copy", int(same))
    print(f"saved work/{os.path.basename(out)} {versions[k_].shape}")

# ------------------------------------------------------------------ 3. checks
# 3a NUTS codes of the release
codes = sig.nuts3_code.dropna(); codes = codes[codes.str.len() == 5]
val21, val16 = float(codes.isin(L21).mean()), float(codes.isin(L16).mean())
print(f"release NUTS3 codes: valid in NUTS 2021 {val21:.4f}, in NUTS 2016 {val16:.4f} ({len(codes)} codes)")
note("release_codes_valid_nuts2021", val21, "S17 and Table S24"); note("release_codes_valid_nuts2016", val16, "S17 and Table S24")
m = sig.dropna(subset=["gcom_id"]).drop_duplicates("gcom_id").merge(t3[["NUTS3"]], left_on="gcom_id", right_index=True, how="inner")
m = m[m.NUTS3.notna()]
dif = m[m.nuts3_code != m.NUTS3]
miss = dif.nuts3_code.isna()
recoded = ~miss & ~dif.NUTS3.isin(L21)       # the 2021-table (NUTS 2016) code no longer exists in NUTS 2021
moved = ~miss & ~recoded                     # a different region
rc = dif[recoded].country_code.value_counts()
print(f"2021 signatories: release code differs from the 2021 table for {len(dif)} of {len(m)} ({len(dif) / len(m) * 100:.1f}%): "
      f"classification changes {int(recoded.sum())} {rc.to_dict()}, different regions {int(moved.sum())}, missing codes {int(miss.sum())}")
note("nuts_codes_differ_share", len(dif) / len(m), "S17"); note("nuts_codes_differ_recoded", int(recoded.sum()), "S17: classification changes")
for c_, n_ in rc.items():
    note(f"nuts_codes_differ_recoded_{c_}", int(n_))
note("nuts_codes_differ_other_region", int(moved.sum()), "S17"); note("nuts_codes_differ_missing", int(miss.sum()), "S17")

# 3b external covariates against the 2021 table (Eurostat GDP not substituted; heating degree days)
gdp = pd.read_csv(GDP_CSV, dtype={"geo": str})
gdp = gdp[(gdp.unit == "EUR_HAB") & gdp.TIME_PERIOD.between(2010, 2018)]
gavg = gdp.groupby("geo").agg(gdp_es=("OBS_VALUE", "mean"))
geo3 = gavg[gavg.index.str.len() == 5]
mb = sig.dropna(subset=["gcom_id"]).merge(t3[["NUTS3", "gdp", "hdd", "hddy"]], left_on="gcom_id", right_index=True, how="inner")   # signatories in both releases
cmp_ = mb.dropna(subset=["NUTS3", "gdp"]).drop_duplicates("NUTS3")[["NUTS3", "gdp"]].merge(geo3, left_on="NUTS3", right_index=True, how="left")
gdp_med = float(np.nanmedian(np.abs(np.log(cmp_.gdp_es / cmp_.gdp))))
hm = mb.dropna(subset=["NUTS3", "hdd", "hddy"])
hv = np.array([hes.get((c, int(y)), np.nan) for c, y in zip(hm.NUTS3, hm.hddy)])
hdd_med = float(np.nanmedian(np.abs(np.log(hv / hm.hdd.values))))
print(f"Eurostat NUTS3 GDP per capita (mean 2010-2018) vs 2021 table: median |log ratio| {gdp_med:.4f} over {int(cmp_.gdp_es.notna().sum())} regions; "
      f"Eurostat HDD vs 2021 table (same region and year): median |log ratio| {hdd_med:.4f}")
note("gdp_eurostat_median_abs_log_ratio", gdp_med, "S17: not substituted"); note("hdd_eurostat_median_abs_log_ratio", hdd_med, "S17 and Table S24")

# 3c sector mapping on inventories with the same id in both releases (Table S24, Panel A)
old["eid"] = old.emission_inventory_id.astype(str)
SH = {"Municipal buildings, equipment/facilities": "Mu", "Tertiary (non municipal) buildings, equipment/facilities": "Te",
      "Residential buildings": "Re", "Transportation": "Tr", "Waste/wastewater": "Wa", "Non-ETS": "NE", "Industry Non-ETS": "NE",
      "ETS (not recommended)": "ETS", "Industry-ETS": "ETS", "Agriculture, Forestry, Fisheries (excluding AFOLU)": "Ag",
      "Buildings, equipment/facilities non allocated": "NA", "Other": "Ot"}
er = pd.concat([pd.read_csv(os.path.join(CSV_DIR, "df2_2D_1.csv"), low_memory=False, dtype={"emission_inventory_id": str}),
                pd.read_csv(os.path.join(CSV_DIR, "df3_3D_1.csv"), low_memory=False, dtype={"emission_inventory_id": str})])
er["s"] = er.aggregated_sector.map(SH)
P = er.pivot_table(index="emission_inventory_id", columns="s", values="emission_measure", aggfunc="sum")
ov = old.set_index("eid").join(P, how="inner")
n_ov = len(ov)


def exact(v21, cols_):
    """2021 inventories with the sector (positive value) and how many are reproduced within 0.01% by the first 7th-release
    category plus any subset of the others."""
    has = v21.notna() & (v21 > 0)
    best = pd.Series(False, index=ov.index)
    for k in range(len(cols_) + 1):
        for combo in itertools.combinations(cols_[1:], k):
            cc = [cols_[0]] + list(combo)
            v7 = ov[cc].fillna(0).sum(axis=1)
            best |= has & (np.abs(v7 - v21) <= TOL * v21)
    return int(has.sum()), int(best.sum())


mp = [("municipal", "Municipal", "Municipal buildings, equipment/facilities", ["Mu"]),
      ("tertiary", "Tertiary", "Tertiary (non municipal) buildings, equipment/facilities", ["Te"]),
      ("residential", "Residential", "Residential buildings; agriculture, forestry, fisheries; buildings, equipment/facilities non allocated", ["Re", "Ag", "NA"]),
      ("manufacturing", "Manufacturing", "Non-ETS; Industry Non-ETS", ["NE"]),
      ("transport", "Transport", "Transportation", ["Tr"]),
      ("waste", "Waste", "Waste/wastewater", ["Wa"])]
mrows = []
for s21, lab, b7, cc in mp:
    v21 = ov[f"{s21}_emissions_tonnes"]
    n_, k_ = exact(v21, cc)
    mrows.append({"2021 sector": lab, "7th-release reporting sectors": b7, "With sector": n_, "Reproduced": k_, "Share": k_ / n_})
    note(f"mapping_{s21}_n", n_); note(f"mapping_{s21}_reproduced_share", k_ / n_, "S17 and Table S24")
tmap = pd.DataFrame(mrows)
n1, k1 = exact(ov["residential_emissions_tonnes"], ["Re"])
ok885 = int((ov.waste_emissions_tonnes.notna() & ov.Wa.notna() & (np.abs(ov.Wa - 0.885 * ov.waste_emissions_tonnes) <= 2e-3 * ov.waste_emissions_tonnes)).sum())
print(f"sector mapping on {n_ov} inventories with the same id:\n{tmap.round(3).to_string(index=False)}")
print(f"residential buildings alone reproduce {k1} of {n1}; waste = 0.885 x 2021 value in {ok885} inventories")
note("inventories_same_id", n_ov, "S17 and Table S24"); note("mapping_residential_buildings_only", k1, f"of {n1}")
note("mapping_waste_0885", ok885, "S17 and Table S24")
# all energy-sector values (direct and indirect emissions) identical for inventories with the same id
t2, _ = load_jrc()
S21 = {"Municipal buildings and facilities": "Mu", "Institutional/tertiary buildings and facilities": "Te", "Residential buildings and facilities": "Re",
       "Manufacturing and construction industries": "Ma", "Transportation": "Tr", "Waste/wastewater": "Wa"}
SHt = dict(SH); SHt.update({"Non-ETS": "Ma", "Industry Non-ETS": "Ma"})
t2x = t2.assign(s=t2.emission_inventory_sector.map(S21), t=np.where(t2.type_of_emissions == "direct_emissions", "dir", "ind"),
                emission_inventory_id=t2.emission_inventory_id.astype(str))
erx = er.assign(s=er.aggregated_sector.map(SHt), t=np.where(er.aggregated_carrier.isin(["Electricity", "District heating and cooling"]), "ind", "dir"))
A2 = t2x.pivot_table(index="emission_inventory_id", columns=["s", "t"], values="emissions", aggfunc="sum")
B7 = erx.pivot_table(index="emission_inventory_id", columns=["s", "t"], values="emission_measure", aggfunc="sum")
ids = A2.index.intersection(B7.index)
A2, B7 = A2.loc[ids], B7.reindex(ids)
okmat = {}
for s_ in ["Mu", "Te", "Re", "Ma", "Tr"]:
    for t_ in ["dir", "ind"]:
        if (s_, t_) in A2:
            a = A2[(s_, t_)]; b = B7[(s_, t_)] if (s_, t_) in B7 else pd.Series(np.nan, index=ids)
            has = a.notna() & (a > 0)
            okmat[(s_, t_)] = np.where(has, (np.abs(b - a) <= TOL * a).astype(float), np.nan)
frac = pd.DataFrame(okmat, index=ids).mean(axis=1)
ident_energy = float((frac == 1).mean())
print(f"inventories with the same id ({len(ids)}): all energy-sector values (direct and indirect) identical in {ident_energy * 100:.1f}%")
note("same_id_energy_values_identical", ident_energy, "Table S24 (direct and indirect emissions of the five energy sectors)")

# 3d analysis samples, the same inventories in both, and the per-capita screen
d21, df21 = load_2021_release()
main21s = release_sample(df21, screen=True)
df7 = prepare_release(inv, inv.NUTS3)
grp7 = set(inv.loc[inv.joint_plan == 1, "GCoM_ID"])
m_all = release_sample(df7)
main7_raw = release_sample(df7, exclude_ids=grp7)
main7 = release_sample(df7, screen=True, exclude_ids=grp7)
print(f"7th release analysis sample: covariates complete and municipal {len(m_all)}; without joint plans {len(main7_raw)}; "
      f"after the per-capita screen {len(main7)} (2021 signatories' covariates {int((main7.covariate_source == 'jrc2021').sum())}, "
      f"new signatories {int((main7.covariate_source == 'external').sum())}); not in 2021: {int(main7.new_inventory.sum())}")
note("joint_plan_inventories_removed", len(m_all) - len(main7_raw), "S17 and Table S24"); note("analysis_sample_7th", len(main7))
note("analysis_inventories_2021_covariates", int((main7.covariate_source == "jrc2021").sum()), "S17 and Table S24")
note("analysis_inventories_new_signatories", int((main7.covariate_source == "external").sum()), "S17 and Table S24")
nw = main7[main7.new_inventory == True]
note("not_in_2021", len(nw)); note("not_in_2021_bei", int((nw.mon == 0).sum())); note("not_in_2021_mei", int((nw.mon == 1).sum()))
note("not_in_2021_median_year", float(nw.inventory_year.median()))
a_all, b_all = keyed(main21s), keyed(main7)
common = sorted(set(a_all.key) & set(b_all.key))
ja = a_all.set_index("key").loc[common]; jb = b_all.set_index("key").loc[common]
EA = ja[SC[:5]].values.astype(float); EB = jb[SC[:5]].values.astype(float)
same_pres = (np.isnan(EA) == np.isnan(EB)).all(1)
rel = np.abs(np.nan_to_num(EB) - np.nan_to_num(EA)) <= TOL * np.maximum(np.abs(np.nan_to_num(EA)), 1e-9)
dl = np.log(jb.total_emissions_reported_tonnes.values) - np.log(ja.total_emissions_reported_tonnes.values)
dp = np.log(jb.population_inventory_year.values) - np.log(ja.population_inventory_year.values)
waste_drop = (ja.sector_set.str[5] == "1").values & (jb.sector_set.str[5] == "0").values
same = dict(n=len(common), energy=float((same_pres & rel.all(1)).mean()), total=float(np.mean(np.abs(dl) < TOL)),
            sector_set=float(np.mean(ja.sector_set.values == jb.sector_set.values)), population=float(np.mean(np.abs(dp) < TOL)),
            waste_dropped=int(waste_drop.sum()), six_2021=int((ja.sc == 6).sum()), six_7th=int((jb.sc == 6).sum()))
print(f"same inventories in both analysis samples: {same['n']}; identical energy-sector values {same['energy'] * 100:.1f}%, totals {same['total'] * 100:.1f}%, "
      f"sector sets {same['sector_set'] * 100:.1f}%, populations {same['population'] * 100:.1f}%; waste reported in 2021 but not in the 7th release "
      f"{same['waste_dropped']}; six-sector inventories {same['six_2021']} -> {same['six_7th']}")
for k_, v_ in same.items():
    note(f"same_inventories_{k_}", v_, "S17, Tables S24-S25")
# waste reported in 2021 but not in the 7th release, across all inventories of the two extracts (source data)
k7 = keyed(inv).set_index("key"); k2 = keyed(old).set_index("key")
both = k2.index.intersection(k7.index)
wl = k2.loc[both, "waste_emissions_tonnes"].notna().values & k7.loc[both, "waste_emissions_tonnes"].isna().values
reg = k2.loc[both, "GCoM_ID"].map(t3.NUTS3).astype(str).str[:4].values
print(f"source data: {int(wl.sum())} inventories report waste in 2021 but not in the 7th release (Andalucia ES61 {int((wl & (reg == 'ES61')).sum())}, "
      f"Catalonia ES51 {int((wl & (reg == 'ES51')).sum())}, Spain {int((wl & (k2.loc[both, 'country_code'].values == 'es')).sum())})")
note("source_waste_lost", int(wl.sum()), "S17 and Table S24"); note("source_waste_lost_es61", int((wl & (reg == "ES61")).sum()))
note("source_waste_lost_es51", int((wl & (reg == "ES51")).sum()))
# per-capita screen
ge7 = main7_raw[main7_raw.gross_error == 1]
ge21 = df21[(df21.gross_error == 1) & (df21.covars_ok == 1) & (df21.nonmunicipal == 0) & df21.nuts3.notna()]
k21i = df21.set_index(["GCoM_ID", "inventory_type", "inventory_year"])
grows = []
for lab, g in [("7th release", ge7), ("2021 dataset", ge21)]:
    for r in g.itertuples():
        row = dict(release=lab, GCoM_ID=r.GCoM_ID, signatory=r.signatory_name, type="BEI" if r.inventory_type.startswith("baseline") else "MEI", year=int(r.inventory_year),
                   total_t=r.total_emissions_reported_tonnes, population=r.population_inventory_year, t_per_capita=r.epc, sectors=r.sector_set)
        kk = (str(r.GCoM_ID), r.inventory_type, int(r.inventory_year))
        if lab == "7th release" and kk in k21i.index:
            q = k21i.loc[[kk]].iloc[0]
            row.update(total_2021_t=q.total_emissions_reported_tonnes, population_2021=q.population_inventory_year,
                       total_ratio_7th_2021=r.total_emissions_reported_tonnes / q.total_emissions_reported_tonnes)
        grows.append(row)
gross = pd.DataFrame(grows)
print("inventories outside 0.01-100 t per capita:\n" + gross.round(4).to_string(index=False))

# ------------------------------------------------------------------ check against the numbers printed in the manuscript
chk = []


def check(where, item, printed, value, digits):
    """Compare a computed value, rounded as printed in the manuscript, with the printed number."""
    v = float(value)
    chk.append(dict(where=where, item=item, manuscript=printed, this_run=round(v, digits), status="OK" if round(v, digits) == printed else "DIFF"))


def num_(name):
    return next(x["value"] for x in nums if x["name"] == name)


g7 = gross[gross.release == "7th release"].set_index("signatory"); g21 = gross[gross.release == "2021 dataset"].set_index("signatory")
for where, item, printed, value, digits in [
        ("S17", "inventories in action plans", 9856, n_ap, 0), ("S17", "inventories in monitoring reports", 3399, n_mr, 0),
        ("S17", "keys with several entries", 1362, multi.sum(), 0), ("S17", "of which with different totals", 635, (tot_diff & multi).sum(), 0),
        ("S17", "distinct keys", 11776, len(inv), 0), ("S17", "inventories in the 2021 source", 8045, len(k21), 0),
        ("S17", "matched by signatory, type and year", 7364, found, 0), ("S17", "inventories with the same id", 4471, n_ov, 0),
        ("S17", "mapping reproduced: municipal, %", 92.0, num_("mapping_municipal_reproduced_share") * 100, 1),
        ("S17", "mapping reproduced: tertiary, %", 88.6, num_("mapping_tertiary_reproduced_share") * 100, 1),
        ("S17", "mapping reproduced: residential, %", 84.8, num_("mapping_residential_reproduced_share") * 100, 1),
        ("S17", "mapping reproduced: manufacturing, %", 86.5, num_("mapping_manufacturing_reproduced_share") * 100, 1),
        ("S17", "mapping reproduced: transport, %", 79.6, num_("mapping_transport_reproduced_share") * 100, 1),
        ("S17", "mapping reproduced: waste, %", 4.1, num_("mapping_waste_reproduced_share") * 100, 1),
        ("S17", "residential buildings alone reproduce", 3032, k1, 0), ("S17", "inventories with residential", 4351, n1, 0),
        ("S17", "zero-emission sector entries", 426, sum(nz.values()), 0), ("S17", "waste 0.885 times the 2021 value", 837, ok885, 0),
        ("Table S24", "same id: identical energy-sector values, %", 55.5, ident_energy * 100, 1),
        ("S17", "joint-plan inventories removed", 160, len(m_all) - len(main7_raw), 0),
        ("S17", "release codes valid in NUTS 2016, %", 95.3, val16 * 100, 1), ("S17", "release codes valid in NUTS 2021, %", 100.0, val21 * 100, 1),
        ("S17", "2021 signatories with a different release code, %", 9.3, len(dif) / len(m) * 100, 1),
        ("S17", "classification changes", 355, recoded.sum(), 0), ("S17", "classification changes in Italy", 257, rc.get("it", 0), 0),
        ("S17", "classification changes in Belgium", 61, rc.get("be", 0), 0), ("S17", "classification changes in Croatia", 36, rc.get("hr", 0), 0),
        ("S17", "different regions", 287, moved.sum(), 0), ("S17", "missing codes", 65, miss.sum(), 0),
        ("S17", "analysis inventories with 2021 covariates", 9764, (main7.covariate_source == "jrc2021").sum(), 0),
        ("S17", "analysis inventories of new signatories", 948, (main7.covariate_source == "external").sum(), 0),
        ("S17", "Eurostat GDP: median absolute difference, %", 9.7, gdp_med * 100, 1),
        ("S17", "NUTS3 agreement on 2021 signatories, %", 96.8, agree * 100, 1),
        ("S17", "density: median absolute difference, %", 1.2, np.nanmedian(np.abs(lr_dens)) * 100, 1),
        ("S17", "heating degree days: median absolute difference, %", 3.3, hdd_med * 100, 1),
        ("S17", "same inventories in both analysis samples", 7011, same["n"], 0),
        ("S17", "same inventories: identical energy-sector values, %", 58.4, same["energy"] * 100, 1),
        ("S17", "same inventories: identical totals, %", 41.3, same["total"] * 100, 1),
        ("S17", "same inventories: identical sector sets, %", 82.2, same["sector_set"] * 100, 1),
        ("S17", "same inventories: identical populations, %", 97.4, same["population"] * 100, 1),
        ("S17", "same inventories: waste reported in 2021 only", 999, same["waste_dropped"], 0),
        ("S17", "six-sector inventories, 2021 values", 774, same["six_2021"], 0), ("S17", "six-sector inventories, 7th values", 307, same["six_7th"], 0),
        ("S17", "source data: waste reported in 2021 only", 1014, wl.sum(), 0),
        ("S17", "source data: of which in Andalucia", 441, (wl & (reg == "ES61")).sum(), 0),
        ("S17", "source data: of which in Catalonia", 400, (wl & (reg == "ES51")).sum(), 0),
        ("S17", "not in 2021", 3701, len(nw), 0), ("S17", "not in 2021: BEIs", 1448, (nw.mon == 0).sum(), 0),
        ("S17", "not in 2021: MEIs", 2253, (nw.mon == 1).sum(), 0), ("S17", "not in 2021: median year", 2015, nw.inventory_year.median(), 0),
        ("S17", "Lingewaard: t per capita", 36721, g7.loc["Lingewaard", "t_per_capita"], 0),
        ("S17", "Lingewaard: total relative to 2021", 45879, g7.loc["Lingewaard", "total_ratio_7th_2021"], 0),
        ("S17", "Orebro: t per capita", 1639, g7.loc[[s for s in g7.index if "rebro" in s][0], "t_per_capita"], 0),
        ("S17", "Orebro: population", 124, g7.loc[[s for s in g7.index if "rebro" in s][0], "population"], 0),
        ("S17", "Orebro: population in 2021", 140000, g7.loc[[s for s in g7.index if "rebro" in s][0], "population_2021"], 0),
        ("S17", "Sant'Omobono Terme: t per capita", 0.0075, g7.loc["Sant'Omobono Terme", "t_per_capita"], 4),
        ("S17", "Ossuccio: total, t", 0.486, g21.loc["Ossuccio", "total_t"], 3),
        ("S17", "Ossuccio: t per capita", 0.0001, g21.loc["Ossuccio", "t_per_capita"], 4)]:
    check(where, item, printed, value, digits)
chk = pd.DataFrame(chk)
chk.to_csv(os.path.join(TAB_DIR, "check_mycovenant7_processing.csv"), index=False)
print(f"\ncheck against the manuscript: {int((chk.status == 'OK').sum())} OK, {int((chk.status == 'DIFF').sum())} DIFF")
if (chk.status == "DIFF").any():
    print(chk[chk.status == "DIFF"].to_string(index=False))

# ------------------------------------------------------------------ outputs
pd.DataFrame(nums).to_csv(os.path.join(TAB_DIR, "mycovenant7_processing_numbers.csv"), index=False)
flow = pd.DataFrame([
    ("Inventories in the 7th release: action plans; monitoring reports", f"{n_ap:,}; {n_mr:,}"),
    ("With a signatory id and a positive six-sector total", f"{next(x['value'] for x in nums if x['name'] == 'inventories_with_id_and_positive_total'):,}"),
    ("One inventory per signatory, type and year (latest entry kept)", f"{len(inv):,}"),
    ("  keys with several entries; of these with different totals", f"{int(multi.sum()):,}; {int((tot_diff & multi).sum()):,}"),
    ("2021 inventories found by signatory, type and year", f"{found:,} of {len(k21):,}"),
    ("Analysis sample: covariates complete, municipal signatories (as in 2021)", f"{len(m_all):,}"),
    ("  without joint plans of several municipalities (number_of_members > 0)", f"{len(main7_raw):,}"),
    ("  without inventories outside 0.01-100 t per capita", f"{len(main7):,}"),
    ("  in both analysis samples / not in 2021", f"{same['n']:,} / {int(main7.new_inventory.sum()):,}"),
], columns=["Step", "Inventories"])
with pd.ExcelWriter(os.path.join(TAB_DIR, "MyCovenant7_Processing.xlsx")) as xw:
    flow.to_excel(xw, sheet_name="Sample_flow", index=False)
    tmap.to_excel(xw, sheet_name="S24_PanelA_mapping", index=False)
    pd.DataFrame(nums).to_excel(xw, sheet_name="Numbers", index=False)
    gross.round(4).to_excel(xw, sheet_name="Per_capita_screen", index=False)
    chk.to_excel(xw, sheet_name="Check_manuscript", index=False)
    hashes.to_excel(xw, sheet_name="Source_files", index=False)


WORDS = {1: "one", 2: "two", 3: "three", 4: "four", 5: "five"}


def pc(x):
    return f"{x * 100:.1f}\\%"


ling = gross[(gross.release == "7th release") & (gross.t_per_capita > EPC_MAX)].sort_values("t_per_capita", ascending=False)
low = gross[(gross.release == "7th release") & (gross.t_per_capita < EPC_MIN)]
n21 = gross[gross.release == "2021 dataset"]


def tex_name(s):
    """Signatory name in LaTeX (the only non-ASCII letter among the listed names is an O with diaeresis)."""
    return str(s).replace("\u00d6", '\\"O')


def screen_txt(r):
    """Description of an inventory removed by the per-capita screen, as in Table S24."""
    head = f"{tex_name(r.signatory)} {r.year} {r.type}"
    ratio = getattr(r, "total_ratio_7th_2021", np.nan)
    pop21 = getattr(r, "population_2021", np.nan)
    if pd.notna(pop21) and r.population / pop21 < 0.01:
        return f"{head} (population {r.population:,.0f}, against {pop21:,.0f} in 2021)"
    pcap = f"{r.t_per_capita:,.0f}" if r.t_per_capita >= 1 else f"{r.t_per_capita:.4f}"
    if pd.notna(ratio) and ratio > 100:
        return f"{head} ({pcap}~t per capita; total {ratio:,.0f} times its 2021 value)"
    if pd.notna(ratio) and abs(ratio - 1) < 1e-6:
        return f"{head} ({pcap}~t per capita, as in 2021)"
    return f"{head} ({pcap}~t per capita)"


items = [screen_txt(r) for r in pd.concat([ling, low]).itertuples()]
excl = items[0] if len(items) == 1 else ", ".join(items[:-1]) + " and " + items[-1]
L = [r"\begin{table}[width=\linewidth,pos=htbp]",
     r"\caption{Processing of the MyCovenant 7th release: sector mapping and data checks}\label{tab:S24}",
     r"\footnotesize", r"\setlength{\tabcolsep}{3pt}",
     rf"\par\medskip\noindent\textit{{Panel A. Sector mapping, checked on {n_ov:,} inventories with the same identifier in both releases}}\par\smallskip",
     r"\noindent\begin{tabularx}{\linewidth}{@{}>{\raggedright\arraybackslash}p{0.20\linewidth}>{\raggedright\arraybackslash}X>{\raggedleft\arraybackslash}p{0.12\linewidth}>{\raggedleft\arraybackslash}p{0.12\linewidth}@{}}",
     r"\toprule", r"2021 sector & 7th-release reporting sectors & With sector & Reproduced \\", r"\midrule"]
for r in tmap.itertuples():
    L.append(f"{r[1]} & {r[2]} & {r[3]:,} & {pc(r[5])} \\\\")
L += [r"Not used & ETS (not recommended); Industry-ETS; Other & & \\", r"\bottomrule", r"\end{tabularx}",
      r"\par\medskip\noindent\textit{Panel B. Data checks}\par\smallskip",
      r"\noindent\begin{tabularx}{\linewidth}{@{}>{\raggedright\arraybackslash}p{0.26\linewidth}>{\raggedright\arraybackslash}X@{}}",
      r"\toprule", r"Item & Finding \\", r"\midrule",
      rf"Release & GCoM -- MyCovenant, 7th Release -- January 2026 (DOI 10.2905/JRC.MDWNXYR); {n_ap:,} inventories in action plans and {n_mr:,} in monitoring reports \\",
      rf"Harmonisation & The 2021 steps (emission factors outside $\pm$50\% of reference values replaced, missing emissions derived from activity data, outliers removed; \citealp{{kona2021global}}) are not described in the release documentation; {pc(ident_energy)} of the {len(ids):,} inventories with the same identifier have identical energy-sector values in both releases \\",
      rf"Same inventories in both analysis samples & {same['n']:,} inventories: identical energy-sector values {pc(same['energy'])}, totals {pc(same['total'])}, sector sets {pc(same['sector_set'])}, population {pc(same['population'])}; {same['waste_dropped']:,} lack the waste reported in 2021, so six-sector donors fall from {same['six_2021']:,} to {same['six_7th']:,} (Table~S25) \\",
      rf"One inventory per signatory, type and year & {len(inv):,} inventories; {int(multi.sum()):,} keys had several entries, {int((tot_diff & multi).sum()):,} with different totals; the latest entry is kept \\",
      rf"Match with the 2021 dataset & {found:,} of the {len(k21):,} inventories of 2021 found by signatory, type and year \\",
      rf"Re-entered inventories & {int(wl.sum()):,} inventories with waste in 2021 appear without waste (Andaluc\'ia {int((wl & (reg == 'ES61')).sum())}, Catalonia {int((wl & (reg == 'ES51')).sum())}) \\",
      rf"NUTS codes & Release codes follow NUTS 2021 (all valid; {pc(val16)} also valid in NUTS 2016); NUTS 2016 is used throughout \\",
      rf"Covariates & 2021 ancillary values for 2021 signatories ({int((main7.covariate_source == 'jrc2021').sum()):,} analysis inventories); for {int((main7.covariate_source == 'external').sum()):,} inventories of new signatories, LAU 2018 population and area, NUTS3 GDP per capita from the 2021 table and NUTS3 heating degree days \\",
      rf"Validation on 2021 signatories & NUTS3 agreement {pc(agree)}; median absolute difference {np.nanmedian(np.abs(lr_dens)) * 100:.1f}\% for density and {hdd_med * 100:.1f}\% for heating degree days \\",
      rf"Exclusions & Joint plans of several municipalities ({len(m_all) - len(main7_raw)} inventories); inventories outside 0.01--100~t per capita ({len(ge7)}): {excl}. The same rule removes {WORDS.get(len(n21), len(n21))} 2021 inventories (Table~S23, second column) \\",
      r"\bottomrule", r"\end{tabularx}",
      rf"\par\smallskip\noindent\begin{{minipage}}{{\linewidth}}\footnotesize Notes: Reproduced: share of 2021 sector values equal to the mapped 7th-release sum within 0.01\%. Agriculture and non-allocated buildings are assigned to the residential sector because this reproduces the 2021 values; residential buildings alone reproduce {k1:,} of the {n1:,}. Waste: in {ok885} inventories the 7th-release value is 0.885 times the 2021 value. Remaining differences mostly combine identical activity with different emissions, consistent with the emission-factor corrections and gap filling of the 2021 harmonisation \citep{{kona2021global}}. Sectors with zero reported emissions are treated as not reported ({sum(nz.values())} entries).\end{{minipage}}",
      r"\end{table}"]
with open(os.path.join(TAB_DIR, "Table_S24.tex"), "w", encoding="utf-8") as f:
    f.write("% Table S24 generated by 23_mycovenant7_extract.py (cas-sc LaTeX)\n" + "\n".join(L) + "\n")
print("saved: work/MyCovenant7_SFA_ready.csv, work/MyCovenant7_SFA_ready_earliest.csv, tables/MyCovenant7_Processing.xlsx, "
      "tables/mycovenant7_processing_numbers.csv, tables/Table_S24.tex")
