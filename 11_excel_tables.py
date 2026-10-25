# -*- coding: utf-8 -*-
"""
11_excel_tables.py - collects the results of the earlier stages into Excel tables (the source of the numbers in the
manuscript tables).

Input : work/*.pkl (stages 01-10)
Output: tables/Main_Tables.xlsx (sheets Table 1-3 = manuscript Tables 2-4),
        tables/Supplementary_Tables.xlsx (Tables S1-S8)
        Derived values (implied differences, shares, ratios) are written as computed values, not as formulas.
"""
from lib import *
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

banner("11 Excel tables (Main_Tables.xlsx, Supplementary_Tables.xlsx)")
data = load_obj("data.pkl"); df, main = data["df"], data["main"]
md = load_obj("model.pkl"); fr = load_obj("frontier.pkl"); tt = load_obj("twotier.pkl"); asym = load_obj("asym.pkl")
bt = load_obj("bootstrap.pkl"); mc = load_obj("montecarlo.pkl"); rp = load_obj("repeated.pkl"); spt = load_obj("spatial.pkl")
bm = load_obj("benchmarks.pkl")

F = "Arial"
HF = Font(name=F, bold=True, size=10); BF = Font(name=F, size=10); TF = Font(name=F, bold=True, size=11); NF = Font(name=F, italic=True, size=9, color="555555")
FILL = PatternFill("solid", start_color="EDEDED"); THIN = Side(style="thin", color="999999")


def title(ws, r, text):
    ws.cell(r, 1, text).font = TF
    return r + 1


def header(ws, r, cols):
    for j, c in enumerate(cols, 1):
        x = ws.cell(r, j, c); x.font = HF; x.fill = FILL; x.border = Border(bottom=THIN)
        x.alignment = Alignment(wrap_text=True, vertical="center")
    return r + 1


def rows(ws, r, recs, fmts):
    for rec in recs:
        for j, v in enumerate(rec, 1):
            if isinstance(v, (np.floating, np.integer)):
                v = v.item()
            if isinstance(v, float) and np.isnan(v):
                v = None
            x = ws.cell(r, j, v); x.font = BF
            if fmts[j - 1] and isinstance(v, (int, float)):
                x.number_format = fmts[j - 1]
        r += 1
    return r


def notes(ws, r, lines):
    for t in lines:
        ws.cell(r, 1, t).font = NF; r += 1
    return r + 1


def widths(ws, w):
    for j, v in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(j)].width = v


# ============================== Main_Tables.xlsx ==============================
labels = {"lnPop": "ln population", "lnGDP": "ln GDP per capita (NUTS3)", "lnDens": "ln population density (2018)", "HDDk": "Heating degree days / 1000",
          "year_c": "Inventory year (centred at 2008)", "sc1": "1 reported sector", "sc2": "2 reported sectors", "sc3": "3 reported sectors",
          "sc4": "4 reported sectors", "sc6": "6 reported sectors", "core5_complete_flag": "Core five sectors complete",
          "co2eq": "CO2-equivalent units", "lca": "LCA emission factors", "mon": "Monitoring inventory"}
specs = md["specs"]
wb = Workbook(); ws = wb.active; ws.title = "Table1_Model"
r = title(ws, 1, "Table 1. Conditional model of reported emissions: coefficients and province-clustered t statistics (dependent variable: ln reported emissions)")
r = header(ws, r, ["Variable", "(1) coef", "(1) t", "(2) coef", "(2) t", "(3) coef", "(3) t", "(4) coef", "(4) t"])
recs = []
for v in ["lnPop", "lnGDP", "lnDens", "HDDk", "year_c", "sc1", "sc2", "sc3", "sc4", "sc6", "core5_complete_flag", "co2eq", "lca", "mon"]:
    rec = [labels[v]]
    for s, sp in enumerate(specs):
        if s == 3 and v in ["lnPop", "lnGDP", "lnDens", "HDDk"]:
            rec += ["spline", None]
        elif v in sp["names"]:
            i = sp["names"].index(v); rec += [sp["b"][i], sp["t"][i]]
        else:
            rec += [None, None]
    recs.append(rec)
recs += [["Country fixed effects", "No", None, "No", None, "Yes", None, "Yes", None],
         ["Natural-spline controls (df = 4)", "No", None, "No", None, "No", None, "Yes", None],
         ["Observations"] + sum([[sp["n"], None] for sp in specs], []), ["R squared"] + sum([[sp["r2"], None] for sp in specs], []),
         ["Residual skewness"] + sum([[sp["skew"], None] for sp in specs], []), ["Coelli M3T"] + sum([[sp["m3t"], None] for sp in specs], [])]
r = rows(ws, r, recs, [None, "0.000", "0.00", "0.000", "0.00", "0.000", "0.00", "0.000", "0.00"]) + 1
r = title(ws, r, "Implied difference in reported emissions relative to five-sector inventories, exp(coef) - 1")
r = header(ws, r, ["Coverage", "Model (3)", "Model (4)"])
r = rows(ws, r, [[labels[v], md["implied"][v][0], md["implied"][v][1]] for v in ["sc1", "sc2", "sc3", "sc4", "sc6"]], [None, "0.0%", "0.0%"])
r = notes(ws, r + 1, [f"Notes: OLS on {len(main):,} inventories of {md['n_muni']:,} municipalities. t statistics clustered by province (NUTS3 identifier; {md['n_clusters']} clusters). Reference coverage: five reported sectors.",
                      "Model (4) replaces the four structural terms by natural cubic splines (df = 4; interior knots at the quartiles, boundary knots at the extremes), so their coefficients are not reported. Municipality-clustered t statistics: sheet Table1_raw.",
                      "Source: 02_conditional_model.py."])
widths(ws, [42, 11, 8, 11, 8, 11, 8, 11, 8])
ws2 = wb.create_sheet("Table1_raw")
r = title(ws2, 1, "Table 1 (raw): coefficients, standard errors and t statistics under province and municipality clustering")
r = header(ws2, r, ["Model", "Variable", "coef", "SE (province)", "t (province)", "SE (municipality)", "t (municipality)"])
recs = []
for s, sp in enumerate(specs):
    for i, v in enumerate(sp["names"]):
        if v.startswith("country_code=") or v.startswith("spline_"):
            continue
        recs.append([f"({s + 1})", labels.get(v, v), sp["b"][i], sp["se"][i], sp["t"][i], sp["se_muni"][i], sp["t_muni"][i]])
rows(ws2, r, recs, [None, None, "0.0000", "0.0000", "0.00", "0.0000", "0.00"]); widths(ws2, [8, 40, 11, 13, 11, 16, 15])

# ---- Table 2 (manuscript Table 3; panels B-D here are its Panels A-C)
ws = wb.create_sheet("Table2_Pairs")
r = title(ws, 1, f"Table 2. Baseline-to-monitoring change in reported emissions by coverage change ({len(rp['pairs']):,} municipalities)")
r = title(ws, r, "Panel A. Mean change in ln reported emissions, all reported sectors and sectors reported in both inventories")
r = header(ws, r, ["Coverage change", "Municipalities", "Median years between inventories", "Mean change, all sectors", "SE", "Mean change, sectors in both", "SE"])
r = rows(ws, r, rp["panelA"][["category", "n", "median_years", "mean_all", "se_all", "mean_both", "se_both"]].values.tolist(),
         [None, "#,##0", "0", "0.000", "0.000", "0.000", "0.000"]) + 1
lab = {"n_rem": "Sectors removed (count)", "n_add": "Sectors added (count)", "dy": "Years between inventories", "dP": "Change in ln population"}
r = title(ws, r, "Panel B. Within-municipality regression: change in ln reported emissions (country fixed effects; province-clustered)")
r = header(ws, r, ["Variable", "coef", "SE", "t"])
r = rows(ws, r, [[lab[x.var], x.coef, x.se, x.t] for x in rp["panelB"].itertuples()] + [["Observations", rp["panelB_n"], None, None], ["R squared", rp["panelB_r2"], None, None]],
         [None, "0.000", "0.000", "0.00"]) + 1
r = title(ws, r, "Panel C. Change in sectors reported in both inventories, relative to constant coverage (country fixed effects; province-clustered)")
r = header(ws, r, ["Variable", "coef", "SE", "t"])
r = rows(ws, r, [[("Years between inventories" if x.var == "dy" else x.var), x.coef, x.se, x.t] for x in rp["panelC"].itertuples()]
         + [["Observations", rp["panelC_n"], None, None], ["R squared", rp["panelC_r2"], None, None]], [None, "0.000", "0.000", "0.00"]) + 1
r = title(ws, r, "Panel D. Single-sector removals: observed change versus change implied by the removed sector's baseline share")
r = header(ws, r, ["Removed sector", "Pairs", "Mean baseline share of removed sector", "Observed mean change in ln total", "Implied change ln(1 - share)", "Observed / implied"])
r = rows(ws, r, rp["panelD"][["sector", "n", "share", "observed", "mechanical", "ratio"]].values.tolist(), [None, "#,##0", "0.0%", "0.000", "0.000", "0.00"])
cc = rp["cc"]
r = notes(ws, r + 1, ["Notes: pairs of a baseline and a later monitoring inventory for the same municipality. Coverage = set of the six sectors with non-missing values. "
                      f"Emission-factor type and unit changes among constant-coverage pairs: {int(cc.ef_change.sum())} and {int(cc.unit_change.sum())}.",
                      "Implied change assumes the removed sector is the only difference between inventories; the ratio above 1 indicates concurrent revisions of the remaining sectors.",
                      "Source: 08_repeated_inventories.py."])
widths(ws, [34, 14, 22, 22, 18, 24, 12])

# ---- Table 3 (manuscript Table 4)
t3 = rp["t3"]
ws = wb.create_sheet("Table3_ConstantCoverage")
r = title(ws, 1, f"Table 3. Annualised change in reported emissions under constant sector coverage ({len(cc):,} municipalities), % per year")
r = header(ws, r, ["Component", "Mean", "95% CI low", "95% CI high", "Median", "Share decreasing", "Selection-weighted mean (IPW)", "Share of total, mean", "Share of total, IPW"])
recs = t3[["component", "mean", "ci_lo", "ci_hi", "median", "share_neg", "ipw_mean", "share_of_total", "share_of_total_ipw"]].values.tolist()
chk = t3.set_index("key")
recs.append(["Check: energy use + carbon intensity - total", chk.loc["A", "mean"] + chk.loc["I", "mean"] - chk.loc["TOT", "mean"], None, None, None, None,
             chk.loc["A", "ipw_mean"] + chk.loc["I", "ipw_mean"] - chk.loc["TOT", "ipw_mean"], None, None])
r = rows(ws, r, recs, [None, "0.00", "0.00", "0.00", "0.00", "0.0%", "0.00", "0.0%", "0.0%"]) + 1
r = title(ws, r, "By country (countries with at least 20 constant-coverage pairs), medians, % per year")
r = header(ws, r, ["Country", "Pairs", "Total", "Per capita", "Energy use", "Carbon intensity"])
r = rows(ws, r, rp["t3_cty"][["country", "pairs", "total", "per_capita", "energy", "intensity"]].values.tolist(), [None, "#,##0", "0.00", "0.00", "0.00", "0.00"])
r = notes(ws, r + 1, ["Notes: change = difference in ln values between monitoring and baseline inventory divided by years elapsed, times 100. ln E = ln(energy use) + ln(carbon intensity) is an identity; means add up, medians do not.",
                      "IPW: weights 1/p from a logit of monitoring submission on ln population, ln GDP per capita, ln density, HDD, baseline year, sector count, core-five completeness, CO2-eq units and the baseline OLS residual (country fixed effects).",
                      "Carbon intensity includes changes in grid and fuel emission factors. Reported energy use is as reported by municipalities. Source: 08_repeated_inventories.py."])
widths(ws, [44, 10, 11, 11, 10, 14, 18, 18, 18])

# ---- Index
wi = wb.create_sheet("Index", 0)
r = title(wi, 1, "Figures and tables written by stages 01-10, with their numbers in the manuscript")
r = header(wi, r, ["Manuscript", "Output file or sheet", "Content", "Stage"])
items = [("Fig. 1", "figures/main/Figure_1", "Shares of inventories by reported-sector count and country (a); reported emissions per capita by sector count (b).", "02"),
         ("Fig. 3", "figures/main/Figure_6", "Reported emissions against the conditional 25th-percentile benchmark (a), gaps by target quantile (b), consistently flagged inventories (c).", "10"),
         ("Fig. S1", "figures/supplementary/Figure_S1", "Locations of the analysed inventories by country.", "01"),
         ("Fig. S2", "figures/supplementary/Figure_S2", "Profile log-likelihood of the inefficiency scale.", "03"),
         ("Fig. S3", "figures/main/Figure_2", "Frontier diagnostics: residual skewness by subsample (a), quantile-quantile plot (b), inefficiency scale after trimming (c).", "03"),
         ("Fig. S4", "figures/supplementary/Figure_S3", "Two-tier frontier, specification (d): components and scales by coverage.", "04"),
         ("Fig. S5", "figures/supplementary/Figure_S4", "Monte Carlo rejection rates (a); estimated excess and shortfall scales across samples (b).", "05, 07"),
         ("Figs. S6, S7", "figures/supplementary/Figure_S5, Figure_S6", "Selection into monitoring; distribution of changes by coverage-change category.", "08"),
         ("Fig. S8", "figures/main/Figure_3", "Coverage change and reported change; single-sector removals.", "08"),
         ("Fig. S9", "figures/main/Figure_5", "Spatial structure of residuals with country and NUTS3 fixed effects.", "09"),
         ("Fig. S11", "figures/main/Figure_4", "Constant-coverage change by country; energy-use and carbon-intensity decomposition.", "08"),
         ("Table 2", "sheets Table1_Model, Table1_raw", "Conditional model of reported emissions.", "02"),
         ("Table 3", "sheet Table2_Pairs, Panels B-D", "Changes between baseline and monitoring inventories by coverage change (Panel A: category means quoted in Section 6.4).", "08"),
         ("Table 4", "sheet Table3_ConstantCoverage", "Constant-coverage change and its decomposition.", "08"),
         ("Tables S1-S8", "Supplementary_Tables.xlsx", "Sample; frontier diagnostics; two-tier model; Monte Carlo; selection; coverage changes; Moran's I; benchmarks.", "01-10")]
r = rows(wi, r, [list(f) for f in items], [None] * 4)
notes(wi, r + 1, ["Sheet and table titles in this workbook use the code's numbering; the first column gives the manuscript number.",
                  "Figures: vector PDF and 300 dpi PNG. The map disclaimer is given in the manuscript captions.",
                  "Inference: standard errors clustered by NUTS3 region ('province' in labels) unless stated; real NUTS3 (2016) codes when USE_REAL_NUTS3 = True.",
                  "Figures 2, 4-6, S10 and S12 and Tables S9-S22 are written by stages 13-22 (see README.md)."])
widths(wi, [14, 44, 110, 8])
wb.save(os.path.join(TAB_DIR, "Main_Tables.xlsx"))

# ========================== Supplementary_Tables.xlsx ==========================
wb = Workbook(); ws = wb.active; ws.title = "S1_Sample"
d = rp["pairs"]
r = title(ws, 1, "Table S1. Sample construction and composition"); r = header(ws, r, ["Item", "Value"])
cshare = main.country_code.value_counts(normalize=True); du = main.degree_urbanisation.value_counts(normalize=True)
recs = [["Inventories in source file", len(df)], ["Signatories in source file", df.GCoM_ID.nunique()], ["Inventories with complete covariates", int(df.covars_ok.sum())],
        ["Non-municipal signatories excluded (entities)", int(df.loc[df.nonmunicipal == 1, "GCoM_ID"].nunique())], ["Non-municipal signatories excluded (inventories)", int(df.nonmunicipal.sum())],
        ["Analysis inventories", len(main)], ["Municipalities", main.GCoM_ID.nunique()], ["Countries", main.country_code.nunique()], ["Province proxies (NUTS3)", main.nuts3.nunique()],
        ["Province proxies, Italy", main[main.country_code == "it"].nuts3.nunique()], ["Province proxies, Spain", main[main.country_code == "es"].nuts3.nunique()],
        ["Province proxies, Belgium", main[main.country_code == "be"].nuts3.nunique()],
        ["Share Italy", cshare.get("it")], ["Share Spain", cshare.get("es")], ["Share Belgium", cshare.get("be")], ["Share Portugal", cshare.get("pt")], ["Share Greece", cshare.get("el")],
        ["Median population", main.population_inventory_year.median()], ["Share with population below 10,000", float((main.population_inventory_year < 10000).mean())],
        ["Degree of urbanisation: cities", du.get(1.0)], ["Degree of urbanisation: towns", du.get(2.0)], ["Degree of urbanisation: rural", du.get(3.0)],
        ["Baseline inventories", int((main.mon == 0).sum())], ["Monitoring inventories", int((main.mon == 1).sum())],
        ["Complete reporters (>= 5 sectors, core five)", int(((main.sc >= 5) & (main.core5_complete_flag == 1)).sum())],
        ["CO2-unit inventories", int((main.co2eq == 0).sum())], ["CO2-equivalent inventories", int((main.co2eq == 1).sum())],
        ["Integer-degree coordinates (excluded from spatial analysis)", int(main.int_coords.sum())],
        ["Population mismatch flag (ratio to 2018 outside 0.5-2)", int(main.pop_mismatch.sum())],
        ["Spatial sample (baseline, non-integer coordinates)", int(((main.mon == 0) & (main.int_coords == 0)).sum())],
        ["Baseline-monitoring pairs", len(d)], ["Constant-coverage pairs", int((d.cat == "same coverage").sum())], ["Median inventory year", main.inventory_year.median()]]
fm = ["#,##0"] * 12 + ["0.0%"] * 5 + ["#,##0", "0.0%", "0.0%", "0.0%", "0.0%"] + ["#,##0"] * 10 + ["0"]   # last row: year
for i_, rec in enumerate(recs):
    ws.cell(r, 1, rec[0]).font = BF
    x = ws.cell(r, 2, float(rec[1]) if rec[1] is not None else None); x.font = BF; x.number_format = fm[i_]; r += 1
notes(ws, r + 1, ["Province: NUTS3 (2016) code from the JRC file when USE_REAL_NUTS3 = True (default); otherwise a GDP-value proxy, (country, NUTS3 GDP per capita value) pairs. Table S9 compares the two.",
                  f"Non-municipal filter (case-insensitive regular expression on signatory names): {NONMUNI_PATTERN}", "Source: 01_data.py."])
widths(ws, [58, 14])

# ---- S2
ws = wb.create_sheet("S2_Frontier"); r = title(ws, 1, "Table S2. Conventional stochastic frontier diagnostics")
r = title(ws, r, "Panel A. Frontier models by one-sided distribution")
r = header(ws, r, ["Model", "Covariates", "sigma_u", "sigma_v", "LR (sigma_u = 0)", "Mean efficiency"])
r = rows(ws, r, fr["dist"][["model", "covariates", "sigma_u", "sigma_v", "LR_sigma_u0", "mean_efficiency"]].values.tolist(), [None, None, "0.000", "0.000", "0.00", "0.000"]) + 1
r = title(ws, r, "Panel B. Residual skewness and frontier scales by subsample (structural + reporting + country FE)")
r = header(ws, r, ["Sample", "n", "Skewness", "Coelli M3T", "Cost sigma_u", "Cost LR", "Production sigma_u", "Production LR"])
r = rows(ws, r, fr["subs"][["sample", "n", "skewness", "M3T", "cost_sigma_u", "cost_LR", "production_sigma_u", "production_LR"]].values.tolist(),
         [None, "#,##0", "0.00", "0.0", "0.000", "0.00", "0.000", "0.00"]) + 1
r = title(ws, r, "Panel C. Trimming sensitivity (cost frontier, half-normal)")
r = header(ws, r, ["Rule", "Share trimmed", "n", "Skewness after trimming", "sigma_u", "LR"])
r = rows(ws, r, fr["trim"][["rule", "share_trimmed", "n", "skewness", "sigma_u", "LR"]].values.tolist(), [None, "0.0%", "#,##0", "0.00", "0.000", "0.00"])
cr = fr["subs"].set_index("sample").loc["Complete reporters (>=5 sectors, core five)"]
notes(ws, r + 1, ["LR tests of sigma_u = 0 have a mixed chi-square null (Kodde-Palm critical values 2.706 at 5% and 5.412 at 1%). A cost frontier adds the one-sided term; the production-type frontier subtracts it.",
                  f"Complete reporters: skewness is small but still negative (M3T about {cr.M3T:.1f}) and a production-type one-sided term remains (LR about {cr.production_LR:.0f}); the low tail is largely, not entirely, a coverage phenomenon.",
                  "Source: 03_frontier_diagnostics.py."])
widths(ws, [62, 36, 10, 10, 14, 14, 16, 14])

# ---- S3
ws = wb.create_sheet("S3_TwoTier"); r = title(ws, 1, "Table S3. Two-tier stochastic frontier (exponential tiers): v + u - w")
r = header(ws, r, ["Specification", "sigma_v", "sigma_u", "sigma_w", "LR (sigma_u = 0) vs shortfall-only", "log-likelihood", "AIC", "BIC", "Mean E[u|e]", "Mean E[w|e]"])
lab = {"a": "(a) structural covariates", "b": "(b) + reporting variables", "c": "(c) + country fixed effects", "d": "(d) + sigma_w depends on reporting", "e": "(e) + sigma_u depends on city type"}
o = tt["out"]
recs = [[lab[k], o[k]["sv"], o[k]["su"], o[k]["sw"], tt["lr"].get(k), o[k]["llf"], o[k]["aic"], o[k]["bic"], o[k]["u"].mean(), o[k]["w"].mean()] for k in "abcde"]
r = rows(ws, r, recs, [None, "0.000", "0.000", "0.000", "0.0", "0.0", "0.0", "0.0", "0.000", "0.000"]) + 1
rd, re_ = o["d"], o["e"]
r = title(ws, r, "Heteroskedasticity of the shortfall scale, specification (d): ln sigma_w = delta0 + z'delta (province-clustered)")
r = header(ws, r, ["Variable", "delta", "SE", "t"])
zl = {"sc_le3": "3 or fewer sectors", "sc4": "4 sectors", "sc6": "6 sectors", "core5_complete_flag": "Core five complete", "mon": "Monitoring inventory"}
r = rows(ws, r, [[zl[c], rd["dw"][j], rd["dw_se"][j], rd["dw"][j] / rd["dw_se"][j]] for j, c in enumerate(ZW_COLS)], [None, "0.000", "0.000", "0.00"]) + 1
r = title(ws, r, "Determinants of the excess scale, specification (e): ln sigma_u = gamma0 + h'gamma (standardised; province-clustered)")
r = header(ws, r, ["Variable", "gamma", "SE", "t"])
r = rows(ws, r, [[c, re_["du"][j], re_["du_se"][j], re_["du"][j] / re_["du_se"][j]] for j, c in enumerate(["ln population", "ln GDP per capita", "ln density", "city or greater city"])],
         [None, "0.000", "0.000", "0.00"]) + 1
r = title(ws, r, "Mean components by number of reported sectors, specification (d)")
r = header(ws, r, ["Sectors", "Mean E[w|e] (shortfall)", "Mean E[u|e] (unexplained excess)"])
r = rows(ws, r, [[int(s), g.w, g.u] for s, g in tt["comp"].iterrows()], ["0", "0.000", "0.000"]) + 1
r = title(ws, r, f"Asymmetry tests by sample, specification (c): LR test of sigma_u = sigma_w (chi-square(1); 1% critical value {CHI2_1_CRIT_1PCT})")
r = header(ws, r, ["Sample", "n", "sigma_v", "sigma_u", "sigma_w", "LR (sigma_u = 0)", "LR (sigma_u = sigma_w)", "Skewness"])
r = rows(ws, r, asym[["sample", "n", "sv", "su", "sw", "LR_su0", "LR_asym", "skew"]].values.tolist(), [None, "#,##0", "0.000", "0.000", "0.000", "0.0", "0.0", "0.00"]) + 1
r = title(ws, r, "Province-cluster bootstrap of specification (c)")
r = header(ws, r, ["Sample", "Replications", "sigma_u 95% interval", "sigma_w 95% interval", "Share sigma_u > sigma_w", "Median LR (asymmetry)", f"Share LR > {CHI2_1_CRIT_1PCT}"])
r = rows(ws, r, [[x.sample, x.B, f"{x.su_lo:.3f}-{x.su_hi:.3f}", f"{x.sw_lo:.3f}-{x.sw_hi:.3f}", x.share_su_gt_sw, x.median_LR, x.share_LR_gt_crit] for x in bt["table"].itertuples()],
         [None, "0", None, None, "0%", "0.0", "0%"])
notes(ws, r + 1, ["LR (sigma_u = 0) is reported for completeness only: under symmetric heavy-tailed noise it rejects with probability one (Table S4), so it is not evidence of an excess tier. Municipality-level scores are not reported.",
                  "Fitting: unrestricted fits start from three points and restricted fits (sigma_u = sigma_w) from four points (asymmetry tests add a grid over the common scale); every start is run to full convergence and the highest likelihood is kept, because the restricted likelihood has local maxima near sigma_v = 0.",
                  "Source: 04_two_tier.py (specifications), 05_asymmetry_tests.py (asymmetry tests), 06_two_tier_bootstrap.py (bootstrap; province-level resampling)."])
widths(ws, [44, 12, 10, 10, 18, 16, 14, 14, 12, 12])

# ---- S4
ws = wb.create_sheet("S4_MonteCarlo")
r = title(ws, 1, f"Table S4. Monte Carlo: two-tier frontier fitted to data with known structure (n = {int(mc.n.iloc[0]):,}; {int(mc.R.iloc[0])} replications)")
r = header(ws, r, ["Data-generating process", "mean sigma_v", "mean sigma_u", "mean sigma_w", "Rejection rate LR(sigma_u = 0) at 1%", "Rejection rate LR(sigma_u = sigma_w) at 1%",
                   "Rejection rate, heteroskedasticity in a random variable", "Mean delta (true 0.8)", "Rejection rate, true heteroskedasticity"])
mc_ = mc.copy(); mc_.loc[mc_.scenario != "wonly_hetero_laplace", ["dw_true", "rej_het_true"]] = np.nan
r = rows(ws, r, mc_[["scenario", "sv", "su", "sw", "rej_LRsu_1pct", "rej_asym_1pct", "rej_het_random_1pct", "dw_true", "rej_het_true"]].values.tolist(),
         [None, "0.000", "0.000", "0.000", "0%", "0%", "0%", "0.00", "0%"])
notes(ws, r + 1, ["Scenarios: normal, laplace, t3, t5, contaminated = symmetric noise only (no one-sided tiers); true2tier = v 0.15, u 0.25, w 0.30; wonly_* = shortfall tier only with normal or Laplace noise; wonly_hetero_laplace = shortfall scale depends on a binary covariate (delta = 0.8).",
                  "Fits as in Table S3 (several starting values, full convergence). Source: 07_monte_carlo.py."])
widths(ws, [24, 12, 12, 12, 20, 22, 26, 14, 20])

# ---- S5
ws = wb.create_sheet("S5_Selection")
r = title(ws, 1, "Table S5. Selection into monitoring: logit of submitting a monitoring inventory (baseline inventories, country fixed effects, province-clustered)")
r = header(ws, r, ["Covariate", "coef", "SE", "t"])
cl_ = {"lnPop": "ln population", "lnGDP": "ln GDP per capita", "lnDens": "ln density", "HDDk": "HDD / 1000", "year_c": "Baseline year (centred)", "sc": "Number of reported sectors",
       "core5_complete_flag": "Core five complete", "co2eq": "CO2-equivalent units", "res": "Baseline OLS residual"}
r = rows(ws, r, [[cl_[x.var], x.coef, x.se, x.t] for x in rp["logit"].itertuples()] + [["Baseline inventories", rp["n_baseline"], None, None],
                                                                                    ["Share submitting monitoring inventory", rp["share_mon"], None, None]],
         [None, "0.000", "0.000", "0.00"])
notes(ws, r + 1, ["Weights 1/p from this model produce the selection-weighted means in Table 3.", "Source: 08_repeated_inventories.py."])
widths(ws, [40, 12, 12, 10])

# ---- S6
ws = wb.create_sheet("S6_CoverageChanges")
r = title(ws, 1, "Table S6. Sector-coverage changes between baseline and monitoring inventories"); r = header(ws, r, ["Sector", "Times removed", "Times added"])
r = rows(ws, r, rp["s6_counts"][["sector", "removed", "added"]].values.tolist(), [None, "#,##0", "#,##0"]) + 1
r = title(ws, r, "Distribution of the change in ln reported emissions by coverage-change category")
r = header(ws, r, ["Category", "n", "p10", "p25", "median", "p75", "p90"])
r = rows(ws, r, rp["s6_dist"][["category", "n", "p10", "p25", "median", "p75", "p90"]].values.tolist(), [None, "#,##0", "0.000", "0.000", "0.000", "0.000", "0.000"])
notes(ws, r + 1, ["Counts refer to sector-level events: a pair removing two sectors contributes two removals. Source: 08_repeated_inventories.py."])
widths(ws, [22, 14, 12, 10, 10, 10, 10])

# ---- S7
ws = wb.create_sheet("S7_Moran")
r = title(ws, 1, f"Table S7. Moran's I of OLS residuals (baseline inventories, n = {spt['n']}; row-standardised k-nearest neighbours; 199 permutations)")
r = header(ws, r, ["k", "I, country FE", "z", "permutation p", "I, NUTS3 FE", "z", "permutation p"])
r = rows(ws, r, spt["moran"][["k", "I_country", "z_country", "p_country", "I_province", "z_province", "p_province"]].values.tolist(),
         ["0", "0.000", "0.0", "0.000", "0.000", "0.0", "0.000"]) + 1
r = title(ws, r, "Within countries (k = 10; residuals re-estimated within each country)")
r = header(ws, r, ["Country", "n", "I, no NUTS3 FE", "z", "I, NUTS3 FE", "z"])
r = rows(ws, r, spt["within"][["country", "n", "I_no_province", "z_no_province", "I_province", "z_province"]].values.tolist(), [None, "#,##0", "0.000", "0.0", "0.000", "0.0"])
_mo = spt["moran"]
_red = (1 - _mo.I_province.abs() / _mo.I_country.abs()) * 100   # reduction of |I| from country to NUTS3 fixed effects
notes(ws, r + 1, ["In absolute value, Moran's I falls by " + ", ".join(f"{v:.1f}%" for v in _red.iloc[:-1]) + f" and {_red.iloc[-1]:.1f}% at k = "
                  + ", ".join(str(int(k)) for k in _mo.k.iloc[:-1]) + f" and {int(_mo.k.iloc[-1])} (Section S7); at the largest k a small residual remains after NUTS3 fixed effects (see p value).",
                  "Source: 09_spatial.py."])
widths(ws, [10, 14, 10, 14, 14, 10, 14])

# ---- S8
ws = wb.create_sheet("S8_Benchmark"); bench = bm["bench"]
r = title(ws, 1, f"Table S8. Peer-quantile benchmarks among complete reporters (n = {bm['n_hq']}) and realised-change references")
r = header(ws, r, ["Target quantile tau", "Aggregate gap (share of reported emissions)", "Share at or below benchmark", "Median gap among inventories above benchmark"])
r = rows(ws, r, [[t, bench[t]["agg"], bench[t]["share_below"], bench[t]["med_above"]] for t in [0.10, 0.25, 0.50]], ["0.00", "0.0%", "0.0%", "0.0%"]) + 1
r = header(ws, r, ["Screening stability (tau = 0.25; 30 province-stratified bootstrap replications)", "Value"])
r = rows(ws, r, [["Consistently flagged (top quartile in >= 90% of replications)", bm["stable"]], ["Share of complete reporters consistently flagged", bm["stable"] / bm["n_hq"]],
                 ["Share with unstable classification (10-90%)", bm["unstable_share"]]], [None, "0.0%"]) + 1
ws.cell(r - 4, 2).number_format = "#,##0"
r = header(ws, r, ["Realised-change reference (constant coverage)", "Annual change (% per year)", "Seven-year cumulative reduction (positive = decline)"])
r = rows(ws, r, [["Top quartile (25th percentile of change)", bm["ref"]["top_quartile"], 1 - np.exp(7 * bm["ref"]["top_quartile"] / 100)],
                 ["Median", bm["ref"]["median"], 1 - np.exp(7 * bm["ref"]["median"] / 100)]], [None, "0.00", "0.0%"])
notes(ws, r + 1, ["Benchmark: conditional tau-quantile of ln reported emissions given structure, reporting controls and country fixed effects (linear programming quantile regression). Gaps are screening measures, not reduction potential.",
                  "Bootstrap resamples provinces with replacement within each country. Source: 10_benchmarks.py."])
widths(ws, [62, 26, 30, 30])
wb.save(os.path.join(TAB_DIR, "Supplementary_Tables.xlsx"))
print("saved: tables/Main_Tables.xlsx, tables/Supplementary_Tables.xlsx")
