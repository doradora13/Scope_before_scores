# -*- coding: utf-8 -*-
"""
18_revision_tables.py - collects the results of stages 13-17 in one Excel workbook and a CSV of key numbers quoted in the text.

Input : work/extra.pkl, scales.pkl, borders_fua.pkl, decomp.pkl, contrast.pkl
Output: tables/Revision_Tables.xlsx (R1-R5 = Tables S9-S13), tables/revision_key_numbers.csv
"""
from lib import *
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

banner("18 Excel tables R1-R5 (Tables S9-S13) and key numbers")
ex = load_obj("extra.pkl"); sc = load_obj("scales.pkl"); bf = load_obj("borders_fua.pkl"); dc = load_obj("decomp.pkl"); ct = load_obj("contrast.pkl")
F = "Arial"; HF = Font(name=F, bold=True, size=10); BF = Font(name=F, size=10); TF = Font(name=F, bold=True, size=11)
NF = Font(name=F, italic=True, size=9, color="555555"); FILL = PatternFill("solid", start_color="EDEDED"); THIN = Side(style="thin", color="999999")


def put_df(ws, r, title, df, fmt="0.000", note=None):
    ws.cell(r, 1, title).font = TF; r += 1
    for j, c in enumerate(df.columns, 1):
        x = ws.cell(r, j, str(c)); x.font = HF; x.fill = FILL; x.border = Border(bottom=THIN); x.alignment = Alignment(wrap_text=True)
    r += 1
    for rec in df.itertuples(index=False):
        for j, v in enumerate(rec, 1):
            if isinstance(v, (np.floating, np.integer)):
                v = v.item()
            if isinstance(v, dict):
                v = ", ".join(f"{a}: {b:.3f}" if isinstance(b, (float, np.floating)) else f"{a}: {b}" for a, b in v.items())
            if isinstance(v, tuple):
                v = str(v)
            if isinstance(v, float) and not np.isfinite(v):
                v = None
            x = ws.cell(r, j, v); x.font = BF
            if isinstance(v, float):
                x.number_format = fmt
        r += 1
    if note:
        ws.cell(r, 1, note).font = NF; r += 1
    return r + 1


def kv(d):
    return pd.DataFrame([(k, v) for k, v in d.items()], columns=["item", "value"])


wb = Workbook(); ws = wb.active; ws.title = "R1 NUTS3"
r = put_df(ws, 1, "Table R1. GDP-value province proxy vs real NUTS3 (2016) codes from the JRC file", kv(ex["concordance"]), fmt="0.000",
           note="1:1 = the proxy group corresponds to exactly one NUTS3 region and vice versa.")
r = put_df(ws, r, "Proxy groups that merge two NUTS3 regions", ex["merged"])
r = put_df(ws, r, "Distinct regions in the analysis sample", kv(ex["levels"]))
ws.column_dimensions["A"].width = 34; ws.column_dimensions["B"].width = 18

ws = wb.create_sheet("R2 scales")
r = put_df(ws, 1, "Table R2a. Residual dependence by fixed-effect unit (baseline inventories, n = %d)" % sc["n"], sc["scales"],
           note="absorbed = share of the variance of country-FE residuals absorbed by the unit's fixed effects; I = Moran's I (k nearest neighbours).")
r = put_df(ws, r, "Table R2b. Real administrative units vs placebo partitions with the same number of groups per country (%d replications)" % int(len(sc["placebo"]) and 200),
           sc["placebo"], note="excess = real - placebo mean; p = share of placebo partitions absorbing at least as much as the real units.")
if sc["vdec"]:
    r = put_df(ws, r, "Table R2c. Nested random-intercept variance shares of country-FE residuals (REML)", kv(sc["vdec"]))
for j in range(1, 18):
    ws.column_dimensions[get_column_letter(j)].width = 14

ws = wb.create_sheet("R3 borders FUA")
r = put_df(ws, 1, "Table R3a. Boundary discontinuity: semivariance of residuals for municipality pairs within 20 km", bf["border"], fmt="0.0000",
           note="d_cross: difference from same-NUTS3 pairs at equal distance (4-km bands) and country; se: dyadic-robust; ci: NUTS2-cluster bootstrap.")
r = put_df(ws, r, "Table R3b. Functional urban areas (GISCO 2018) vs NUTS3", kv(bf["fua"]))
r = put_df(ws, r, "Table R3c. Joint inventories and identical per-capita values", kv(bf["joint"]))
for j in range(1, 16):
    ws.column_dimensions[get_column_letter(j)].width = 15

ws = wb.create_sheet("R4 decomposition")
r = put_df(ws, 1, "Table R4a. Exact decomposition of constant-coverage change (% per year)", dc["three"],
           note="TOT = A + CIen + NE; CIen = STR + EFi + EFd (LMDI, pairs with positive direct and indirect parts). share_* relative to TOT.")
r = put_df(ws, r, "Table R4b. Emission-factor regime: implied grid factor vs JRC national factors (NEEFE)", kv(dc["ef"]))
r = put_df(ws, r, "Table R4c. Emission-factor regime by country", dc["ef_cty"])
r = put_df(ws, r, "Table R4d. Signs of estimated (not measured) activity data", kv(dc["est"]))
r = put_df(ws, r, "Table R4e. Influence of the most influential NUTS2 region on the pooled decomposition", kv(dc["infl"]))
r = put_df(ws, r, "Table R4f. Leave-one-NUTS2-out pooled decomposition (10 lowest energy-use shares)", dc["loo"].head(10))
for j in range(1, 18):
    ws.column_dimensions[get_column_letter(j)].width = 14

ws = wb.create_sheet("R5 screening")
r = put_df(ws, 1, "Table R5a. Baseline-to-monitoring change with and without coverage screening", kv(ct["ts"]))
r = put_df(ws, r, "Table R5b. Annualised change by coverage-change category (% per year)", ct["cat_tab"])
r = put_df(ws, r, "Table R5c. Coverage composition of apparently low-emitting inventories", kv(ct["xs"]))
ws.column_dimensions["A"].width = 30; ws.column_dimensions["B"].width = 16
out = os.path.join(TAB_DIR, "Revision_Tables.xlsx"); wb.save(out)

# ---- key numbers quoted in the text
s = sc["scales"].set_index("level"); p = sc["placebo"]
pN3k = p[(p["sample"] == "all") & (p.level == "NUTS3") & (p.placebo == "kmeans")].iloc[0]
pN3v = p[(p["sample"] == "all") & (p.level == "NUTS3") & (p.placebo == "voronoi")].iloc[0]
b = bf["border"].set_index("sample"); fu = bf["fua"]; t3 = dc["three"].set_index("sample"); e_ = dc["ef"]
key = [
    ("proxy groups mapping 1:1 to a real NUTS3 region", ex["concordance"]["proxies_one_to_one"]),
    ("inventories in 1:1 proxy groups (%)", 100 * ex["concordance"]["share_one_to_one"]),
    ("Moran I k=10: country / NUTS1 / NUTS2 / NUTS3 FE", "%.3f / %.3f / %.3f / %.3f" % tuple(s.I10[["country", "NUTS1", "NUTS2", "NUTS3"]])),
    ("Moran I k=10 after k-means placebo FE (mean, 95% range)", "%.3f [%.3f, %.3f]" % (pN3k.placebo_I10_mean, pN3k.placebo_I10_lo, pN3k.placebo_I10_hi)),
    ("variance absorbed (%): NUTS3 vs k-means vs voronoi placebo", "%.1f vs %.1f vs %.1f" % (100 * pN3k.real_absorbed, 100 * pN3k.placebo_mean, 100 * pN3v.placebo_mean)),
    ("NUTS3 boundary within NUTS2: semivariance difference (%)", 100 * b.loc["all", "rel_N3"]),
    ("NUTS2 boundary: semivariance difference (%)", 100 * b.loc["all", "rel_N2"]),
    ("FUA subsample size; F(FUA|NUTS3) p; F(NUTS3|FUA) p", "%d; %.3f; %.2g" % (fu["n_inside"], fu["p_fua_given_n3"], fu["p_n3_given_fua"])),
    ("constant coverage: energy use / energy C-intensity / non-energy (% per year)",
     "%.2f / %.2f / %.2f" % tuple(t3.loc["all constant-coverage", ["A", "CIen", "NE"]])),
    ("C-intensity LMDI: electricity share / grid EF / fuel EF (% per year)", "%.2f / %.2f / %.2f" % tuple(t3.loc["all constant-coverage", ["STR", "EFi", "EFd"]])),
    ("most influential NUTS2 region; its share of pairs (%); its share of the pooled energy-use term (%)",
     "%s; %.1f; %.0f" % (dc["infl"]["region"], 100 * dc["infl"]["share_of_pairs"], 100 * dc["infl"]["contrib_A"])),
    ("energy-use share of the pooled change without that region (%)", 100 * dc["infl"]["share_A_without"]),
    ("implied grid EF identical in BEI and MEI (%)", 100 * e_["share_grid_ef_identical"]),
    ("slope of dln implied grid EF on dln national factor", e_["slope"]),
    ("mean annual change: all pairs vs constant coverage (%)", "%.2f vs %.2f" % (ct["ts"]["ann_all"], ct["ts"]["ann_same"])),
    ("share of summed decline from coverage-change pairs (%)", 100 * ct["ts"]["share_sum_from_changed"]),
]
pd.DataFrame(key, columns=["quantity", "value"]).to_csv(os.path.join(TAB_DIR, "revision_key_numbers.csv"), index=False, encoding="utf-8-sig")
print(pd.DataFrame(key, columns=["quantity", "value"]).to_string(index=False))
print(f"saved: {out}, tables/revision_key_numbers.csv")
